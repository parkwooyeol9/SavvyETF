"""ETF universe snapshots from the SavvyDB Refinitiv masters (one file per listing market).

Input : ~/GitHub/SavvyDB/ETF_DB/ETF_Master_*.xlsx (Refinitiv/Datastream sheets, refreshed by hand)
Output: 00_exports/etf_universe/ETF_Universe_<시장>_<YYYYMMDD>.xlsx   (analyst workbook per market)
        webapp/src/data/etfUniverse/<market>.json + index.json       (bundled, served by /api/etf-universe)

Every market goes through the same taxonomy so the tab can switch markets without
re-deriving anything client-side:
  자산군 · 투자지역 · 업종(GICS, 주식형만) · 전략/테마 · 구분(일반/레버리지/인버스) · 운용방식(패시브/액티브)

Keyword rules are ported from webapp/src/lib/etfDb.ts (Korean GICS keywords) and
webapp/src/lib/etfDbUsClassify.ts (English theme rules) so the new tab keeps the old taxonomy.

Usage:
  python3 -I 03_ETF/etf_universe.py                    # all markets
  python3 -I 03_ETF/etf_universe.py --markets kr us    # subset
"""
from __future__ import annotations

import argparse
import json
import re
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

REPO = Path(__file__).resolve().parents[1]
SRC_DIR = Path.home() / "GitHub" / "SavvyDB" / "ETF_DB"
XLSX_DIR = REPO / "00_exports" / "etf_universe"
JSON_DIR = REPO / "webapp" / "src" / "data" / "etfUniverse"

VALID_FTYPE = {"Equity", "Bond", "Alternatives", "Commodity", "Mixed Assets", "Money Market", "Other"}

# ─────────────────────────────── taxonomy ───────────────────────────────

GEO_KO = {
    "United States of America": "미국", "North America": "북미", "Canada": "캐나다",
    "Global": "글로벌", "Global Ex US": "글로벌(미국 제외)", "Global Emerging Markets": "신흥국",
    "Frontier Markets": "프런티어", "BRIC": "신흥국", "Europe": "유럽", "EuroZone": "유로존",
    "Europe exc UK": "유럽", "Europe exc Switzerland": "유럽", "European Misc": "유럽", "Eastern Europe": "유럽",
    "Nordic": "유럽", "United Kingdom": "영국", "Germany": "독일", "France": "프랑스", "Italy": "이탈리아",
    "Spain": "스페인", "Switzerland": "스위스", "Netherlands": "네덜란드", "Belgium": "벨기에",
    "Austria": "오스트리아", "Denmark": "덴마크", "Finland": "핀란드", "Sweden": "스웨덴", "Norway": "노르웨이",
    "Ireland": "아일랜드", "Greece": "그리스", "Poland": "폴란드", "Turkey": "튀르키예", "Russia": "러시아",
    "China": "중국", "Greater China": "중화권", "Hong Kong": "홍콩", "Taiwan": "대만", "Japan": "일본",
    "Korea": "한국", "India": "인도", "Pakistan": "파키스탄", "Vietnam": "베트남", "Indonesia": "인도네시아",
    "Malaysia": "말레이시아", "Philippines": "필리핀", "Thailand": "태국", "Singapore": "싱가포르",
    "ASEAN": "아세안", "Asia Pacific": "아시아·태평양", "Asia Pacific ex Japan": "아시아·태평양",
    "Asia (ex-Japan)": "아시아·태평양", "Far East exc Japan": "아시아·태평양", "Australia": "호주",
    "New Zealand": "뉴질랜드", "Latin America": "중남미", "Brazil": "브라질", "Mexico": "멕시코",
    "Chile": "칠레", "Colombia": "콜롬비아", "Peru": "페루", "Argentina": "아르헨티나",
    "Middle East": "중동", "GCC": "중동", "Saudi Arabia": "사우디", "United Arab Emirates": "중동",
    "Qatar": "중동", "Kuwait": "중동", "Israel": "이스라엘", "Africa": "아프리카", "South Africa": "남아공",
}

# English-name fallbacks when Refinitiv has no geographic focus
GEO_NAME_RULES = [
    ("대만", r"\btaiwan\b"), ("인도", r"\bindia|nifty|sensex"), ("일본", r"\bjapan|nikkei|topix|jpx"),
    ("홍콩", r"\bhong\s*kong|hang\s*seng|\bhsi\b|hscei"), ("중국", r"\bchina|\bcsi\b|\bsse\b|szse|chinext|star\s*50|\bcni\b|a50"),
    ("한국", r"\bkorea|kospi|kosdaq"), ("베트남", r"vietnam"), ("유로존", r"euro\s*stoxx|\bemu\b|eurozone"),
    ("유럽", r"\beurope|stoxx"), ("영국", r"\buk\b|ftse\s*(100|250)"), ("독일", r"\bdax\b|germany"),
    ("신흥국", r"emerging|\bem\b"), ("미국", r"s&p\s*500|nasdaq|\bus\b|u\.s\.|russell|dow\s*jones"),
    ("글로벌", r"world|global|acwi|all[\s-]*country"),
]

SECTOR_KO = {
    "Information Technology": "IT", "Healthcare": "헬스케어", "Biotechnology": "헬스케어",
    "Financials": "금융", "Industrials": "산업재", "Consumer Discretionary": "경기소비재",
    "Consumer Staples": "필수소비재", "Energy": "에너지", "Materials": "소재",
    "Gold & Precious Metals": "소재", "Communication Services": "커뮤니케이션", "Utilities": "유틸리티",
}
THEME_SECTOR = {
    "Alternative Energy": "유틸리티", "EV and Future Mobility": "경기소비재", "Infrastructure": "산업재",
    "Natural Resources": "소재", "Agribusiness": "소재", "Water": "유틸리티",
}

# webapp/src/lib/etfDb.ts GICS_SECTOR_RULES (Korean) + etfDbUsClassify.ts sector words (English)
GICS_RULES = [
    ("IT", r"반도체|HBM|소프트웨어|테크|(?-i:\bIT\b)|semiconductor|technology|\btech\b|software|cloud|cyber|\bai\b|artificial intelligence|robot|internet|chip"),
    ("헬스케어", r"헬스케어|의료|제약|바이오|health|biotech|pharma|medical|genomic"),
    ("금융", r"금융|은행|증권|보험|financ|bank|insurance|fintech"),
    ("에너지", r"에너지|원유|석유|천연가스|energy|oil|gas|uranium|nuclear|원자력|원전"),
    ("소재", r"소재|철강|구리|리튬|화학|금광|materials?|mining|miner|steel|chemical|lithium|copper|gold miner"),
    ("산업재", r"산업재|방산|우주|항공|조선|해운|건설|인프라|운송|기계|industrial|defen[cs]e|aerospace|infrastructure|transport"),
    ("경기소비재", r"자동차|자율주행|화장품|유통|리테일|게임|엔터|소비재|consumer discretionary|retail|\bev\b|electric vehicle|luxury|명품"),
    ("필수소비재", r"필수소비|consumer staples"),
    ("커뮤니케이션", r"통신|인터넷|미디어|콘텐츠|communication|media|telecom"),
    ("유틸리티", r"유틸|전력|태양광|풍력|그리드|utilit|clean energy|solar|wind|renewable|power"),
    ("부동산", r"리츠|부동산|REIT|real estate|property"),
]

BENCH_RE = re.compile(
    r"s&p\s*500|sp\s*500|nasdaq[\s-]*100|나스닥\s*100|나스닥100|dow jones industrial|다우존스30|russell\s*(1000|2000|3000)|"
    r"total\s*(stock\s*)?market|msci\s*(world|acwi|emerging|eafe|europe|emu|japan|usa|china|india|korea|taiwan|brazil|em\b|all country|ac )|"
    r"ftse\s*(100|250|all[\s-]*world|developed|emerging|mib|china|japan|taiwan)|stoxx\s*(europe\s*)?(50|600)|euro\s*stoxx\s*50|유로스탁스50|"
    r"\bdax\b|cac\s*40|\bsmi\b|\baex\b|ibex|nikkei\s*225|니케이\s*225|topix|jpx[\s-]*nikkei\s*400|csi\s*(300|500|800|1000|a50|a500)|"
    r"\bsse\s*50|star\s*50|chinext|hang\s*seng\s*(index|china enterprises|tech)?|\bhscei\b|\bhsi\b|항셍|코스피\s*200|코스피200|kospi\s*200|"
    r"코스닥\s*150|kosdaq\s*150|krx\s*300|nifty\s*50|니프티\s*50|sensex|s&p/asx|\btsx\b|vn30|s&p500|"
    r"core\s*(s&p|msci|ftse)|\bcrsp\b|broad\s*market|extended\s*market",
    re.I,
)
FACTOR_RE = re.compile(r"\bvalue\b|growth|momentum|quality|min(imum)?\s*vol|low\s*vol|equal[\s-]*weight|multi[\s-]*factor|fundamental|smart\s*beta|"
                       r"가치|성장|모멘텀|퀄리티|저변동|동일가중|밸류", re.I)
SMID_RE = re.compile(r"small|mid[\s-]*cap|sm&mid|\bsmid\b|micro[\s-]*cap|중소형|소형주|중형주", re.I)
ESG_RE = re.compile(r"\besg\b|\bsri\b|\bpab\b|paris|climate|sustainab|low\s*carbon|\bctb\b|net\s*zero|탄소|ESG", re.I)
COVERED_RE = re.compile(r"covered\s*call|buy[\s-]*write|premium\s*income|option\s*income|equity\s*premium|buffer|defined\s*outcome|"
                        r"accelerated|autocall|커버드콜|프리미엄|버퍼|옵션", re.I)
DIV_RE = re.compile(r"dividend|high\s*div|aristocrat|\bincome\b|배당|인컴|SCHD", re.I)
SEMI_RE = re.compile(r"semiconductor|\bsemis?\b|\bchips?\b|\bsox\b|\bhbm\b|memory|반도체|TSMC|필라델피아", re.I)
TECH_RE = re.compile(r"technology|\btech\b|artificial\s*intelligence|\bai\b|robot|cloud|software|cyber|internet|digital|innovation|"
                     r"fintech|blockchain|quantum|metaverse|테크|로봇|소프트웨어|클라우드|양자|플랫폼|휴머노이드|빅테크", re.I)
# Case-sensitive: "AI" must not hit Daiwa/Huatai, and "TOP10" means a tech basket only in Korean product names
TECH_CS_RE = re.compile(r"(?<![A-Za-z])AI(?![a-z])")
KR_TOP_RE = re.compile(r"TOP\s?\d")

LEV_RE = re.compile(r"\b(2x|3x|4x|1\.5x|2\.5x|ultra(pro)?|leverag\w*|levrgd|bull\b|daily\s+target\s+2x)|레버리지", re.I)
INV_RE = re.compile(r"\b(inverse|bear\b|ivr\.?|-1x|-2x|-3x|short(?!\s*(-|\s)?(term|duration|maturity|dated|treasury|bond|t-?bill|vol)))|인버스", re.I)
ACTIVE_RE = re.compile(r"\bactive\b|액티브", re.I)

ISSUER_PREFIX = [
    ("BRJ iShares", "iShares"), ("iShares", "iShares"), ("IShares", "iShares"), ("Vanguard", "Vanguard"),
    ("SPDR", "SPDR"), ("State Street", "SPDR"), ("Invesco", "Invesco"), ("Xtrackers", "Xtrackers"), ("Xtr ", "Xtrackers"),
    ("Amundi", "Amundi"), ("AMUNDI", "Amundi"), ("Lyxor", "Amundi"), ("Multi Units", "Amundi"), ("UBS", "UBS"),
    ("BNPP", "BNP Paribas"), ("BNP", "BNP Paribas"), ("HSBC", "HSBC"), ("JPM", "JPMorgan"), ("JPMorgan", "JPMorgan"),
    ("First Trust", "First Trust"), ("Global X", "Global X"), ("GlobalX", "Global X"), ("Franklin", "Franklin Templeton"),
    ("WisdomTree", "WisdomTree"), ("VanEck", "VanEck"), ("Fineco", "Fineco"), ("AXA SPDB", "AXA SPDB"), ("AXA", "AXA IM"),
    ("Goldman Sachs", "Goldman Sachs"), ("L&G", "L&G"), ("Legal & General", "L&G"), ("Market Access", "Market Access"),
    ("IndexIQ", "IndexIQ"), ("OSSIAM", "Ossiam"), ("Ossiam", "Ossiam"),
    ("Nomura", "Nomura (NEXT FUNDS)"), ("NOMU", "Nomura (NEXT FUNDS)"), ("STATE SPDR", "SPDR"), ("NEXT FUNDS", "Nomura (NEXT FUNDS)"), ("NFS", "Nomura (NEXT FUNDS)"),
    ("Daiwa", "Daiwa"), ("Amova", "Amova (구 Nikko AM)"), ("Nikko", "Amova (구 Nikko AM)"), ("MUAM", "MUFG (MAXIS)"),
    ("Mitsubishi UFJ", "MUFG (MAXIS)"), ("Nochu", "Nochu NZAM"), ("AM-One", "Asset Management One"),
    ("Simplex", "Simplex"), ("SMDS", "Sumitomo Mitsui DS"), ("SMT", "Sumitomo Mitsui Trust"), ("Rakuten", "Rakuten"),
    ("Nissay", "Nissay"), ("Japan Physical", "Mitsubishi UFJ Trust"),
    ("ICBC CSOP", "CSOP"), ("CSOP", "CSOP"), ("Hang Seng", "Hang Seng Inv."), ("Tracker Fund", "State Street (TraHK)"),
    ("E Fund", "E Fund"), ("ChinaAMC", "ChinaAMC"), ("Premia", "Premia"), ("Ping An", "Ping An"), ("Ping-An", "Ping An"),
    ("Value Partners", "Value Partners"), ("Samsung", "Samsung"), ("Bosera", "Bosera"), ("China Universal", "China Universal"),
    ("China Southern", "China Southern"), ("Guotai", "Guotai"), ("Fullgoal", "Fullgoal"), ("Hua An", "Hua An"),
    ("GF ", "GF Fund"), ("Yinhua", "Yinhua"), ("Hwabao", "Hwabao WP"), ("Harvest", "Harvest"), ("CMF", "China Merchants"),
    ("Penghua", "Penghua"), ("Huatai", "Huatai-PineBridge"), ("Tianhong", "Tianhong"), ("ICBCUBS", "ICBC Credit Suisse/UBS"),
    ("CCB", "CCB Principal"), ("Dacheng", "Dacheng"), ("TaiKang", "Taikang"), ("Taikang", "Taikang"), ("HFT", "HFT"),
    ("Founder Fubon", "Founder Fubon"), ("China Life", "China Life AMC"), ("Tibet Eastmoney", "Eastmoney"),
    ("CPIC", "CPIC"), ("PengYang", "PengYang"), ("Harfor", "Harfor"), ("SWS MU", "SWS MU"),
]
US_ISSUER = [
    ("BlackRock", "iShares"), ("SSGA", "SPDR"), ("Vanguard", "Vanguard"), ("Invesco", "Invesco"), ("Schwab", "Schwab"),
    ("First Trust", "First Trust"), ("ProShare", "ProShares"), ("Rafferty", "Direxion"), ("Tidal", "Tidal"),
    ("Innovator", "Innovator"), ("Empowered", "Alpha Architect"), ("Global X", "Global X"), ("WisdomTree", "WisdomTree"),
    ("JP Morgan", "JPMorgan"), ("Fidelity", "Fidelity"), ("Van Eck", "VanEck"), ("PGIM", "PGIM"), ("Dimensional", "Dimensional"),
    ("American Century", "American Century (Avantis)"), ("Goldman", "Goldman Sachs"), ("Pacer", "Pacer"), ("DBX", "Xtrackers"),
    ("T.Rowe", "T. Rowe Price"), ("Northern Trust", "Northern Trust (FlexShares)"), ("Franklin", "Franklin Templeton"),
    ("Pacific Investment", "PIMCO"), ("Krane", "KraneShares"), ("Capital Research", "Capital Group"), ("ARK", "ARK"),
    ("Allianz", "Allianz (AllianzIM)"), ("GraniteShares", "GraniteShares"), ("Roundhill", "Roundhill"), ("Simplify", "Simplify"),
    ("Amplify", "Amplify"), ("Nuveen", "Nuveen"), ("Alps", "ALPS"), ("YieldMax", "YieldMax"), ("Neos", "NEOS"),
    ("Morgan Stanley", "Morgan Stanley (Eaton Vance)"), ("Janus", "Janus Henderson"), ("Grayscale", "Grayscale"),
]

KR_MANAGER = {
    "KODEX": "삼성자산운용", "TIGER": "미래에셋자산운용", "RISE": "KB자산운용", "ACE": "한국투자신탁운용",
    "PLUS": "한화자산운용", "SOL": "신한자산운용", "KIWOOM": "키움투자자산운용", "HANARO": "NH-Amundi자산운용",
    "1Q": "하나자산운용", "WON": "우리자산운용", "HK": "흥국자산운용", "BNK": "BNK자산운용", "마이티": "DB자산운용",
    "파워": "교보악사자산운용", "UNICORN": "현대자산운용", "IBK": "IBK자산운용", "DAISHIN": "대신자산운용",
    "TIME": "타임폴리오자산운용", "에셋플러스": "에셋플러스자산운용", "KoAct": "삼성액티브자산운용",
    "TRUSTON": "트러스톤자산운용", "MIDAS": "마이다스에셋",
}
KR_OVERSEAS = {"북미", "아시아", "글로벌", "해외", "유럽", "선진국", "신흥국", "남미", "프런티어"}
KR_REGION_RULES = [
    ("대만", ["대만", "타이완", "TSMC"]), ("베트남", ["베트남"]), ("인도네시아", ["인도네시아"]), ("필리핀", ["필리핀"]),
    ("인도", ["인도", "Nifty", "니프티"]), ("일본", ["일본", "니케이", "TOPIX", "Nikkei", "JPX"]),
    ("중국", ["중국", "차이나", "China", "항셍", "Hang Seng", "홍콩", "CSI", "과창판", "심천", "상해", "커촹"]),
    ("독일", ["독일", "DAX"]), ("유럽", ["유럽", "유로", "Euro", "STOXX"]), ("멕시코", ["멕시코"]),
    ("중남미", ["브라질", "라틴", "남미"]), ("신흥국", ["신흥", "이머징", "Emerging"]),
    ("미국", ["미국", "S&P 500", "S&P500", "나스닥", "NASDAQ", "Nasdaq", "다우", "Dow", "필라델피아", "테슬라", "엔비디아", "애플", "팔란티어", "빅테크"]),
    ("글로벌", ["글로벌", "Global", "선진", "World", "ACWI"]),
]
KR_MARKET_FALLBACK = {"북미": "미국", "아시아": "아시아·태평양", "유럽": "유럽", "남미": "중남미", "신흥국": "신흥국",
                      "선진국": "글로벌", "글로벌": "글로벌", "해외": "글로벌", "프런티어": "프런티어"}

EXCHANGE = {
    "K": "NYSE Arca", "O": "Nasdaq", "N": "NYSE", "A": "NYSE American", "Z": "Cboe BZX", "PK": "OTC",
    "DE": "Xetra", "PA": "Euronext Paris", "L": "LSE", "AS": "Euronext Amsterdam", "S": "SIX", "MI": "Borsa Italiana",
    "LU": "Luxembourg", "BR": "Euronext Brussels", "MC": "BME", "VI": "Vienna", "ST": "Stockholm", "OL": "Oslo",
    "HK": "HKEX", "SS": "상하이", "SZ": "선전", "T": "도쿄", "NG": "나고야", "TO": "TSX",
}


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def text(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.startswith("Unable to collect") or s in ("NULL", "Index is not available on Lipper Database"):
        return None
    return s


def issuer_of(name: str, company: str | None = None) -> str:
    if company:
        for key, label in US_ISSUER:
            if key.lower() in company.lower():
                return label
    for key, label in ISSUER_PREFIX:
        if name.lower().startswith(key.lower()):
            return label
    if company:
        return re.sub(r",?\s*(LLC|LP|Inc\.?|Ltd|Co|Corp(oration)?|Management|Advisors?|Investments?|Asset|Capital|Fund|Company)\b\.?", "", company).strip(" ,") or company
    return name.split()[0] if name else "기타"


def leverage_of(name: str, sector: str | None) -> str:
    sec = sector or ""
    if "Dedicated Short Bias" in sec or INV_RE.search(name):
        return "인버스"
    if "Equity Leveraged" in sec or LEV_RE.search(name):
        return "레버리지"
    return "일반"


def asset_of(ftype: str | None, sector: str | None, name: str) -> str:
    sec = sector or ""
    if "Cryptocurrency" in sec or re.search(r"bitcoin|ether(eum)?\b|crypto|solana|\bxrp\b|비트코인|이더리움", name, re.I):
        return "가상자산"
    if re.search(r"currency|달러선물|엔선물|\bfx\b|\byen\b|\beuro\s+currency", name, re.I) and ftype != "Bond":
        return "통화"
    if "Real Estate" in sec or re.search(r"\breit|리츠|부동산|real\s*estate|property", name, re.I):
        return "리츠·부동산"
    if ftype == "Equity":
        return "주식"
    if ftype == "Money Market" or re.search(r"money\s*market|머니마켓|CD금리|KOFR|SOFR|t-?bill|초단기|cash", name, re.I):
        return "단기자금"
    if ftype == "Bond" or sec.startswith(("Bond", "Target Maturity Bond", "Loan")):
        return "채권"
    if ftype == "Commodity" or sec.startswith("Commodity"):
        return "원자재"
    if ftype == "Mixed Assets" or sec.startswith(("Mixed Asset", "Target Maturity MA")):
        return "혼합"
    if ftype == "Alternatives":
        if re.search(r"gold|silver|oil|crude|gas|commodit|copper|carbon|금|은선물|원유|구리|탄소배출권", name, re.I):
            return "원자재"
        if re.search(r"treasur|bond|채권|국채", name, re.I):
            return "채권"
        if ("Equity" in sec or "Relative Value" in sec or leverage_of(name, sec) != "일반" or COVERED_RE.search(name)
                or re.search(r"equity|stock|dividend|s&p|nasdaq|russell|dow\b|premium|msci|나스닥|코스닥|코스피|200", name, re.I)):
            return "주식"
        return "대체·기타"
    # No Refinitiv fund type (new listings): infer from the name, equity by default
    if re.search(r"회사채|특수채|국채|국고채|통안채|WGBI|treasur|\bbond", name, re.I):
        return "채권"
    if re.search(r"CD\d?년?금리|금리|머니마켓", name):
        return "단기자금"
    if re.search(r"gold|silver|oil|crude|carbon|금현물|원유|탄소배출권", name, re.I):
        return "원자재"
    return "주식" if ftype is None else "대체·기타"


def gics_of(asset: str, sector: str | None, name: str, bm: str | None) -> str | None:
    if asset == "리츠·부동산":
        return "부동산"
    if asset != "주식":
        return None
    sec = sector or ""
    m = re.match(r"Equity Sector (.+?)(?: (?:Asia Pacific|Global|US|Japan|Europe|UK|Australia))?$", sec)
    if m and m.group(1) in SECTOR_KO:
        return SECTOR_KO[m.group(1)]
    m = re.match(r"Equity Theme - (.+)$", sec)
    if m and m.group(1) in THEME_SECTOR:
        return THEME_SECTOR[m.group(1)]
    s = f"{name} {bm or ''}"
    for label, pat in GICS_RULES:
        if re.search(pat, s, re.I):
            return label
    return "시장 전체(분산)"


def strategy_of(asset: str, name: str, bm: str | None, sector: str | None) -> str:
    s = f"{name} {bm or ''}"
    sec = sector or ""
    if asset == "주식":
        if COVERED_RE.search(name):
            return "커버드콜·옵션"
        if DIV_RE.search(name) or "Income" in sec:
            return "배당"
        if SEMI_RE.search(name):
            return "반도체"
        if TECH_RE.search(name) or TECH_CS_RE.search(name) or (re.search(r"[가-힣]", name) and KR_TOP_RE.search(name)):
            return "테크·AI"
        if ESG_RE.search(name):
            return "ESG"
        if FACTOR_RE.search(name):
            return "팩터·스마트베타"
        if SMID_RE.search(name) or "Sm&Mid" in sec or "Small & Mid" in sec:
            return "중소형"
        if BENCH_RE.search(s) and not re.search(r"sector|theme|섹터|테마", s, re.I) and not sec.startswith(("Equity Sector", "Equity Theme")):
            return "대표지수"
        if sec.startswith(("Equity Sector", "Equity Theme")):
            return "섹터·테마"
        if re.match(r"Equity [A-Z]", sec) and not re.search(r"[가-힣]{2,}", name):
            return "광범위 시장"
        return "섹터·테마"
    if asset == "채권":
        if re.search(r"high\s*yield|하이일드|junk", s, re.I) or "High Yield" in sec:
            return "하이일드"
        if re.search(r"inflation|tips|물가", s, re.I) or "Inflation" in sec:
            return "물가연동채"
        if re.search(r"muni|municipal", s, re.I) or "Municipal" in sec:
            return "지방채"
        if re.search(r"mortgage|\bmbs\b|securitized|abs\b", s, re.I) or "Mortgage" in sec:
            return "MBS·ABS"
        if re.search(r"target|maturity|ibond|bulletshares|20\d\d|만기", s, re.I) or sec.startswith("Target Maturity"):
            return "만기매칭"
        if re.search(r"emerging|\bem\b|신흥", s, re.I) or "Emerging" in sec:
            return "신흥국채"
        if re.search(r"corporate|corp\b|credit|회사채|크레딧|투자등급", s, re.I) or "Corporates" in sec:
            return "회사채"
        if re.search(r"short|1-3|0-1|0-3|단기|floating|float", s, re.I) or "Short Term" in sec:
            return "단기채"
        if re.search(r"treasur|government|govt|gilt|bund|jgb|국채|국고채|통안", s, re.I) or "Government" in sec:
            return "국채"
        return "종합·기타채권"
    if asset == "원자재":
        if re.search(r"gold|silver|platinum|palladium|precious|금|은선물|은현물|귀금속", s, re.I) or "Precious" in sec:
            return "금·귀금속"
        if re.search(r"oil|crude|brent|wti|gas|energy|원유|천연가스", s, re.I) or "Energy" in sec:
            return "원유·에너지"
        if re.search(r"copper|alumin|nickel|zinc|metal|구리|산업금속", s, re.I) or "Industrial Metals" in sec:
            return "산업금속"
        if re.search(r"agri|grain|corn|wheat|soy|sugar|농산물", s, re.I) or "Agriculture" in sec:
            return "농산물"
        return "원자재종합"
    return asset


# ─────────────────────────────── loaders ───────────────────────────────

def _sheet_rows(path: Path, sheet: str, start: int):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    return list(ws.iter_rows(min_row=start, values_only=True))


def _date(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and 19000101 < v < 21000101:
        s = str(int(v))
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    if isinstance(v, str):
        m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})$", v.strip())
        if m:
            return f"{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    return None


def _common(name, ric, geo, ftype, sector, bm, home_region):
    lev = leverage_of(name, sector)
    asset = asset_of(ftype, sector, name)
    region = GEO_KO.get(geo or "")
    if not region:
        for label, pat in GEO_NAME_RULES:
            if re.search(pat, f"{name} {bm or ''}", re.I):
                region = label
                break
    suffix = ric.split(".", 1)[1] if "." in ric else ""
    return {
        "ric": ric,
        "n": name,
        "reg": region or home_region,
        "ac": asset,
        "sec": gics_of(asset, sector, name, bm),
        "st": strategy_of(asset, name, bm, sector),
        "lev": lev,
        "act": 1 if ACTIVE_RE.search(name) else 0,
        "bm": bm,
        "ex": EXCHANGE.get(suffix, suffix or None),
        "rf": sector,
    }


def load_refinitiv(path: Path, sheet: str, lay: dict, home_region: str, extra: dict | None = None):
    out, dropped = [], defaultdict(int)
    for r in _sheet_rows(path, sheet, 3):
        ric = r[lay["ric"]]
        if not isinstance(ric, str) or not ric.strip():
            continue
        ric = ric.strip()
        if "^" in ric:
            dropped["상장폐지(RIC ^)"] += 1
            continue
        if ric.endswith(".PK"):
            dropped["OTC(.PK)"] += 1
            continue
        ftype = text(r[lay["ftype"]])
        if ftype not in VALID_FTYPE:
            dropped["Refinitiv 펀드유형 없음(ETF 아님·미수집)"] += 1
            continue
        aum = num(r[lay["aum"]])
        if not aum or aum <= 0:
            dropped["시가총액 없음"] += 1
            continue
        name = text(r[lay["name"]])
        if not name or re.match(r"Japan Shell ETF|.*Securities Holdings", name):
            # Refinitiv placeholder ("Japan Shell ETF 057") or issuer-company name → Datastream product name
            name = text(r[lay.get("name2", lay["name"])]) or name or ric
        # "iShares Inc - iShares MSCI Taiwan ETF" / "Vanguard Funds PLC - Vanguard ..." → product name only
        name = re.sub(r"^[^-]{2,60}?\b(Trust|Inc|Funds?|PLC|plc|ICAV|SICAV|Series \d+|Company)\b[^-]{0,20}\s+-\s+", "", name)
        sector, geo, bm = text(r[lay["sector"]]), text(r[lay["geo"]]), text(r[lay["bm"]])
        rec = _common(name, ric, geo, ftype, sector, bm, home_region)
        company = text(r[lay["company"]]) if "company" in lay else None
        ticker = text(r[lay["ticker"]]) if "ticker" in lay else None
        adv = num(r[lay["adv"]])
        if adv is not None and lay.get("adv_fx") == "mv":
            local = num(r[lay["mv_local"]])
            adv = adv * aum / local if local else None  # quote currency → USD via the row's own MV ratio
        if adv is not None and lay.get("adv_div"):
            adv = adv / lay["adv_div"]
        rec.update({
            "t": ticker or ric.split(".")[0],
            "iss": issuer_of(name, company),
            "ccy": text(r[lay["ccy"]]) if "ccy" in lay else lay.get("ccy_const"),
            "ld": _date(r[lay["listed"]]) if "listed" in lay else None,
            "ter": num(r[lay["ter"]]),
            "aum": aum,
            "adv": adv,
            "r1w": num(r[lay["r1w"]]) if "r1w" in lay else None,
            "r1m": num(r[lay["r1m"]]),
            "r3m": num(r[lay["r3m"]]),
            "r6m": num(r[lay["r6m"]]) if "r6m" in lay else None,
            "r1y": num(r[lay["r1y"]]),
            "vol": num(r[lay["vol"]]),
            "shp": num(r[lay["shp"]]),
        })
        if "div" in lay:
            rec["dv"] = num(r[lay["div"]])
        if extra:
            rec.update(extra)
        out.append(rec)
    return out, dict(dropped)


# Column layouts (0-based) of the Refinitiv sheets; aum is always X(MV)~U$ in USD millions.
LAY_US = dict(ric=10, ticker=27, name=11, name2=4, geo=12, ter=13, bm=14, shp=15, r1y=17, vol=18, adv=19, adv_div=1e6,
              ftype=21, sector=22, company=23, div=24, listed=1, aum=6, r1w=7, r1m=8, r3m=9, ccy_const="USD")
LAY_EU = dict(ric=9, name=10, geo=11, ter=12, bm=13, shp=14, vol=17, adv=18, adv_fx="mv", mv_local=2, adv_div=1e6,
              ftype=20, sector=21, ccy=22, listed=23, aum=3, r1w=5, r1m=6, r3m=7, r1y=8)
LAY_HKCN = dict(LAY_EU, aum=4)
LAY_JP = dict(ric=9, name=10, name2=2, geo=11, ter=12, bm=13, shp=14, vol=17, adv=18, adv_div=1e6,
              ftype=20, sector=21, aum=4, r1m=5, r3m=6, r6m=7, r1y=8, ccy_const="JPY")


def load_kr(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Update_260813"]
    stamp = str(ws.cell(1, 2).value or "").replace("Last Update : ", "")
    out = []
    for r in ws.iter_rows(min_row=14, values_only=True):
        if not r[0] or not r[1]:
            continue
        name, idx, mk = r[1], text(r[3]), r[8]
        ftype, geo, sector = text(r[48]), text(r[46]), text(r[47])
        lev_raw = (r[6] or "").strip()
        lev = "일반" if lev_raw in ("일반", "") else ("레버리지" if "레버" in lev_raw else "인버스")
        fx_future = "달러" in name and "선물" in name
        asset = asset_of(ftype, sector, name)
        if fx_future:
            asset = "통화"
        if re.search(r"채권혼합|TDF|TRF|혼합", name):
            asset = "혼합"
        if mk in KR_OVERSEAS or (mk == "주식외" and fx_future):
            scope = "해외"
            region = next((reg for reg, kws in KR_REGION_RULES if any(k in name for k in kws)), None) \
                or next((reg for reg, kws in KR_REGION_RULES if any(k in (idx or "") for k in kws)), None) \
                or GEO_KO.get(geo or "") or KR_MARKET_FALLBACK.get(mk, "글로벌")
        elif mk == "국내&해외":
            scope, region = "국내·해외 혼합", "국내·해외 혼합"
        else:
            scope, region = "국내", "한국"
        close, close_1y = num(r[20]), num(r[25])
        aum_bil = num(r[9])
        brand = name.split()[0]
        sec = gics_of(asset, sector, name, idx)
        if asset == "주식" and scope == "국내" and sec == "시장 전체(분산)" and not BENCH_RE.search(f"{name} {idx}"):
            sec = "시장 전체(분산)"
        out.append({
            "t": str(r[0]).lstrip("A"),
            "ric": text(r[45]),
            "n": name,
            "iss": KR_MANAGER.get(brand, brand),
            "reg": region,
            "scope": scope,
            "mk": mk,
            "ac": asset,
            "sec": sec,
            "st": strategy_of(asset, name, idx, sector),
            "lev": lev,
            "act": 1 if ACTIVE_RE.search(name) else 0,
            "hd": 1 if re.search(r"\(H\)|H\)|\(H|환헤지", name) else 0,
            "bm": idx,
            "ex": "KRX",
            "ccy": "KRW",
            "ld": _date(r[4]),
            "ter": num(r[5]),
            "aum": aum_bil * 10 if aum_bil is not None else None,  # KRW BIL → 억원
            "adv": num(r[19]) * 10 if num(r[19]) is not None else None,
            "r1m": num(r[33]), "r3m": num(r[34]), "r6m": num(r[35]),
            "r1y": (close / close_1y - 1) * 100 if close and close_1y else None,
            "gap": num(r[17]),
            "te": num(r[53]),
            "mdd": num(r[54]),
            "shp": num(r[50]),
            "dv": num(r[56]),
            "tax": text(r[7]),
            "rf": sector,
        })
    return out, stamp, {}


def file_stamp(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M")


def load_market(key: str):
    m = MARKETS[key]
    path = SRC_DIR / m["file"]
    if key == "kr":
        return load_kr(path)
    recs, dropped = [], defaultdict(int)
    for sheet, extra in m["sheets"]:
        part, d = load_refinitiv(path, sheet, m["layout"], m["home"], extra)
        recs += part
        for k, v in d.items():
            dropped[k] += v
    # Ireland and Luxembourg masters can list the same UCITS line (same ticker, name, AUM).
    if len(m["sheets"]) > 1:
        seen = set()
        unique = []
        for rec in recs:
            key = (rec.get("t"), rec.get("n"), rec.get("aum"))
            if key in seen:
                dropped["시트 간 중복(동일 티커·이름·순자산)"] += 1
                continue
            seen.add(key)
            unique.append(rec)
        recs = unique
    return recs, f"{file_stamp(path)} (파일 저장 시각)", dict(dropped)


MARKETS = {
    "kr": dict(label="한국", file="ETF_Master_한국(QT).xlsx", unit="억원", adv_unit="억원"),
    "us": dict(label="미국", file="ETF_Master_미국.xlsx", sheets=[("ETF_Master", None)], layout=LAY_US, home="미국",
               unit="$M", adv_unit="$M"),
    "eu": dict(label="유럽", file="ETF_Master_유럽 (상장 거래소 기준).xlsx",
               sheets=[("ETF_Master_아일랜드", {"dom": "아일랜드"}), ("ETF_Master_룩셈부르그", {"dom": "룩셈부르크"})],
               layout=LAY_EU, home="유럽", unit="$M", adv_unit="$M"),
    "jp": dict(label="일본", file="ETF_Master_일본.xlsx", sheets=[("ETF_Master_일본", None)], layout=LAY_JP, home="일본",
               unit="$M", adv_unit="¥M"),
    "hk": dict(label="홍콩", file="ETF_Master_홍콩_중국.xlsx", sheets=[("ETF_Master_홍콩", None)], layout=LAY_HKCN, home="홍콩",
               unit="$M", adv_unit="$M"),
    "cn": dict(label="중국", file="ETF_Master_홍콩_중국.xlsx", sheets=[("ETF_Master_중국", None)], layout=LAY_HKCN, home="중국",
               unit="$M", adv_unit="$M"),
}

# ─────────────────────────────── outputs ───────────────────────────────

FIELDS = [  # key, header, width, number format
    ("t", "티커/코드", 10, None), ("n", "종목명", 44, None), ("iss", "운용사", 20, None), ("ex", "상장거래소", 13, None),
    ("dom", "설정국", 10, None), ("scope", "국내/해외", 11, None), ("reg", "투자지역", 13, None), ("ac", "자산군", 10, None),
    ("sec", "업종(GICS)", 13, None), ("st", "전략/테마", 14, None), ("lev", "구분", 9, None), ("act", "운용방식", 9, None),
    ("hd", "환헤지", 9, None), ("bm", "기초지수(BM)", 42, None), ("mk", "기초시장", 11, None), ("ccy", "통화", 7, None),
    ("ld", "상장일", 11, None), ("ter", "총보수(%)", 9, "0.00##"), ("aum", "순자산", 12, "#,##0"),
    ("adv", "일평균 거래대금", 12, "#,##0.0"), ("r1w", "1W(%)", 8, "0.0;[Red]-0.0"), ("r1m", "1M(%)", 8, "0.0;[Red]-0.0"),
    ("r3m", "3M(%)", 8, "0.0;[Red]-0.0"), ("r6m", "6M(%)", 8, "0.0;[Red]-0.0"), ("r1y", "1Y(%)", 8, "0.0;[Red]-0.0"),
    ("vol", "변동성 1Y(%)", 9, "0.0"), ("gap", "괴리율(%)", 8, "0.00;[Red]-0.00"), ("te", "추적오차 1Y(%)", 9, "0.00"),
    ("mdd", "MDD 1Y(%)", 9, "0.0;[Red]-0.0"), ("shp", "샤프 1Y", 8, "0.00"), ("dv", "최근 분배금", 10, "#,##0.00"),
    ("tax", "과세유형", 24, None), ("ric", "RIC", 13, None), ("rf", "Refinitiv 분류", 34, None),
]
RET_KEYS = ("r1w", "r1m", "r3m", "r6m", "r1y")

NAVY = "1F3864"
HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_FONT = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="맑은 고딕", size=10)
BOLD_FONT = Font(name="맑은 고딕", bold=True, size=10)
TITLE_FONT = Font(name="맑은 고딕", bold=True, size=14, color=NAVY)
SUB_FONT = Font(name="맑은 고딕", size=9, color="666666")
SEC_FONT = Font(name="맑은 고딕", bold=True, size=11, color=NAVY)
BAND = PatternFill("solid", fgColor="F2F5FA")
BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))
CENTER = {"t", "ex", "dom", "scope", "reg", "ac", "sec", "st", "lev", "act", "hd", "mk", "ccy", "ld"}


def display(key, v):
    if key == "act":
        return "액티브" if v else "패시브"
    if key == "hd":
        return "환헤지(H)" if v else "환노출"
    return v


def present_fields(recs):
    return [f for f in FIELDS if any(r.get(f[0]) is not None for r in recs)]


def write_table(ws, recs, fields, title, subtitle, unit, adv_unit):
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font, ws["A2"].font = TITLE_FONT, SUB_FONT
    hr = 4
    for j, (key, head, width, _) in enumerate(fields, 1):
        label = f"{head}({unit})" if key == "aum" else f"{head}({adv_unit})" if key == "adv" else head
        c = ws.cell(hr, j, label)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[hr].height = 30
    for i, rec in enumerate(recs, hr + 1):
        for j, (key, _, _, fmt) in enumerate(fields, 1):
            c = ws.cell(i, j, display(key, rec.get(key)))
            c.font, c.border = BODY_FONT, BORDER
            if fmt:
                c.number_format = fmt
            if key in CENTER:
                c.alignment = Alignment(horizontal="center")
    last = hr + len(recs)
    ws.freeze_panes = ws.cell(hr + 1, 3)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(fields))}{max(last, hr)}"
    if not recs:
        return
    col = {f[0]: get_column_letter(j) for j, f in enumerate(fields, 1)}
    ws.conditional_formatting.add(f"{col['aum']}{hr+1}:{col['aum']}{last}",
                                  DataBarRule(start_type="min", end_type="max", color="8EA9DB"))
    for k in RET_KEYS:
        if k in col:
            ws.conditional_formatting.add(
                f"{col[k]}{hr+1}:{col[k]}{last}",
                ColorScaleRule(start_type="percentile", start_value=5, start_color="F8696B", mid_type="num", mid_value=0,
                               mid_color="FFFFFF", end_type="percentile", end_value=95, end_color="63BE7B"))


def group(recs, key):
    g = defaultdict(list)
    for r in recs:
        g[display(key, r.get(key)) or "—"].append(r)
    total = sum(r["aum"] or 0 for r in recs) or 1
    rows = []
    for label, items in g.items():
        aum = sum(x["aum"] or 0 for x in items)
        fees = [x["ter"] for x in items if x.get("ter") is not None]
        r1y = [(x["r1y"], x["aum"] or 0) for x in items if x.get("r1y") is not None]
        w = sum(a for _, a in r1y)
        top = max(items, key=lambda x: x["aum"] or 0)
        rows.append({
            "label": label, "count": len(items), "aum": aum, "share": aum / total * 100,
            "ter": sum(fees) / len(fees) if fees else None,
            "r1y": sum(v * a for v, a in r1y) / w if w else None,
            "top": top["n"],
        })
    rows.sort(key=lambda x: -x["aum"])
    return rows


def summary_block(ws, row, title, recs, key, unit, limit=None):
    ws.cell(row, 1, title).font = SEC_FONT
    row += 1
    heads = ["구분", "상품 수", f"순자산({unit})", "비중(%)", "평균 총보수(%)", "1Y 수익률(%, AUM가중)", "최대 상품"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row, j, h)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    rows = group(recs, key)
    shown = rows[:limit] if limit else rows
    start = row + 1
    for g in shown:
        row += 1
        for j, v in enumerate([g["label"], g["count"], g["aum"], g["share"], g["ter"], g["r1y"], g["top"]], 1):
            c = ws.cell(row, j, v)
            c.font, c.border = BODY_FONT, BORDER
        ws.cell(row, 3).number_format = "#,##0"
        ws.cell(row, 4).number_format = "0.0"
        ws.cell(row, 5).number_format = "0.00"
        ws.cell(row, 6).number_format = "0.0;[Red]-0.0"
    if shown:
        ws.conditional_formatting.add(f"C{start}:C{row}", DataBarRule(start_type="num", start_value=0, end_type="max", color="8EA9DB"))
    if limit and len(rows) > limit:
        row += 1
        rest = rows[limit:]
        ws.cell(row, 1, f"기타 {len(rest)}개").font = SUB_FONT
        ws.cell(row, 2, sum(g["count"] for g in rest)).font = SUB_FONT
        ws.cell(row, 3, sum(g["aum"] for g in rest)).number_format = "#,##0"
        ws.cell(row, 4, sum(g["share"] for g in rest)).number_format = "0.0"
    row += 1
    total = sum(r["aum"] or 0 for r in recs)
    for j, v in enumerate(["합계", len(recs), total, 100.0], 1):
        c = ws.cell(row, j, v)
        c.font, c.fill = BOLD_FONT, BAND
    ws.cell(row, 3).number_format = "#,##0"
    ws.cell(row, 4).number_format = "0.0"
    return row + 2


def build_workbook(key, recs, stamp, dropped, out_path):
    m = MARKETS[key]
    unit, adv_unit, label = m["unit"], m["adv_unit"], m["label"]
    recs = sorted(recs, key=lambda x: -(x["aum"] or 0))
    fields = present_fields(recs)
    total = sum(r["aum"] or 0 for r in recs)
    passive = [r for r in recs if not r["act"]]
    plain = [r for r in passive if r["lev"] == "일반"]
    total_txt = f"{total/1e4:,.1f}조원" if key == "kr" else f"${total/1e6:,.2f}T" if total >= 1e6 else f"${total/1e3:,.1f}B"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{label} 상장 ETF 유니버스 요약"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"원본: SavvyDB/ETF_DB/{m['file']} · 기준 {stamp} · {len(recs):,}개 · 순자산 {total_txt} · "
                f"패시브 {len(passive):,}개({sum(r['aum'] or 0 for r in passive)/total*100:.1f}%)")
    ws["A2"].font = SUB_FONT
    row = 4
    if key == "kr":
        row = summary_block(ws, row, "① 국내 / 해외", recs, "scope", unit)
    row = summary_block(ws, row, "② 자산군", recs, "ac", unit)
    row = summary_block(ws, row, "③ 투자지역", recs, "reg", unit, limit=20)
    eq = [r for r in recs if r["ac"] == "주식"]
    row = summary_block(ws, row, "④ 업종(GICS) — 주식형", eq, "sec", unit)
    row = summary_block(ws, row, "⑤ 전략/테마 — 주식형", eq, "st", unit)
    bond = [r for r in recs if r["ac"] == "채권"]
    if bond:
        row = summary_block(ws, row, "⑥ 채권 세부", bond, "st", unit)
    row = summary_block(ws, row, "⑦ 운용사", recs, "iss", unit, limit=20)
    row = summary_block(ws, row, "⑧ 패시브 / 액티브", recs, "act", unit)
    row = summary_block(ws, row, "⑨ 일반 / 레버리지 / 인버스", recs, "lev", unit)
    if key == "eu":
        row = summary_block(ws, row, "⑩ 설정국(도미사일)", recs, "dom", unit)
        row = summary_block(ws, row, "⑪ 상장거래소(대표 RIC)", recs, "ex", unit)
    if key in ("us", "cn"):
        row = summary_block(ws, row, "⑩ 상장거래소", recs, "ex", unit)

    ws.cell(row, 1, "순자산 상위 30").font = SEC_FONT
    row += 1
    top_fields = [f for f in fields if f[0] in ("n", "t", "iss", "reg", "ac", "st", "ter", "aum", "r3m", "r1y")]
    order = ["n", "t", "iss", "reg", "ac", "st", "aum", "ter", "r3m", "r1y"]
    top_fields.sort(key=lambda f: order.index(f[0]))
    for j, f in enumerate(top_fields, 1):
        c = ws.cell(row, j, f"{f[1]}({unit})" if f[0] == "aum" else f[1])
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center")
    for rec in recs[:30]:
        row += 1
        for j, f in enumerate(top_fields, 1):
            c = ws.cell(row, j, display(f[0], rec.get(f[0])))
            c.font, c.border = BODY_FONT, BORDER
            if f[3]:
                c.number_format = f[3]
    for col, w in zip("ABCDEFGHIJ", [40, 10, 16, 12, 22, 16, 40, 10, 10, 10]):
        ws.column_dimensions[col].width = w

    sub = f"순자산 내림차순 · 필터로 자산군/지역/전략/구분 선택 · 기준 {stamp}"
    write_table(wb.create_sheet("전체"), recs, fields, f"{label} 상장 ETF 전체", f"{len(recs):,}개 · {sub}", unit, adv_unit)
    write_table(wb.create_sheet("패시브(일반형)"), plain, fields, f"{label} 상장 패시브 ETF — 레버리지·인버스 제외",
                f"{len(plain):,}개 · {sub}", unit, adv_unit)
    if key == "kr":
        ovs = [r for r in passive if r["scope"] == "해외"]
        write_table(wb.create_sheet("해외 패시브"), ovs, fields, "국내 상장 해외 패시브 ETF", f"{len(ovs):,}개 · {sub}", unit, adv_unit)
        mixed = [r for r in passive if r["scope"] == "국내·해외 혼합"]
        write_table(wb.create_sheet("국내·해외 혼합"), mixed, fields, "참고: 국내·해외 자산 혼합형 패시브 ETF",
                    f"{len(mixed)}개 (채권혼합·TDF·TRF 등)", unit, adv_unit)
    act = [r for r in recs if r["act"]]
    write_table(wb.create_sheet("액티브"), act, fields, f"{label} 상장 액티브 ETF (종목명 기준)", f"{len(act):,}개 · {sub}", unit, adv_unit)

    note = wb.create_sheet("기준·주석")
    note["A1"] = "기준 및 주석"
    note["A1"].font = TITLE_FONT
    rows = notes_for(key, stamp, dropped)
    for i, (k, v) in enumerate(rows, 3):
        note.cell(i, 1, k).font = BOLD_FONT
        note.cell(i, 2, v).font = BODY_FONT
        note.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    note.column_dimensions["A"].width = 18
    note.column_dimensions["B"].width = 120
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def notes_for(key, stamp, dropped):
    m = MARKETS[key]
    sheets = "Update_260813" if key == "kr" else ", ".join(s for s, _ in m["sheets"])
    rows = [
        ("데이터 원천", f"SavvyDB/ETF_DB/{m['file']} · 시트 {sheets} (Refinitiv/Datastream) · 기준 {stamp}"),
        ("운용방식", "종목명에 '액티브'/'Active'가 들어가면 액티브, 아니면 패시브 (한국 마스터 파일과 같은 규칙)"),
        ("구분", "한국: 원본 'ETF레버리지,인버스구분' 컬럼. 그 외: Refinitiv 분류(Equity Leveraged / Dedicated Short Bias)와 종목명(2X·3X·Ultra·Leveraged·Inverse·Bear·Short) 키워드"),
        ("자산군", "Refinitiv 펀드유형(TR.FundType)·분류를 기본으로, 가상자산·통화·리츠·단기자금은 종목명 키워드로 따로 뺌. Alternatives는 기초자산(주식·원자재·채권)으로 다시 배정"),
        ("투자지역", "Refinitiv 지역(TR.FundGeographicFocus)을 한글로 옮김. 비어 있으면 종목명·BM 키워드, 그래도 없으면 상장 시장"),
        ("업종(GICS)", "주식형만. Refinitiv 'Equity Sector/Theme' 분류 우선, 없으면 종목명 키워드(SavvyETF 기존 DB 탭 규칙 이식). 섹터 키워드가 없으면 '시장 전체(분산)'"),
        ("전략/테마", "종목명·BM 키워드 순서: 커버드콜·옵션 → 배당 → 반도체 → 테크·AI → ESG → 팩터 → 중소형 → 대표지수 → 섹터·테마/광범위 시장. 채권·원자재는 세부 유형"),
        ("순자산", "한국: 원본 AUM(KRW BIL) × 10 = 억원. 그 외: Datastream X(MV)~U$ = 백만 달러"),
        ("거래대금", "한국: 억원. 미국: TR.AvgDailyValTraded20D(USD) ÷ 1e6. 유럽·홍콩·중국: 거래통화 금액을 같은 행의 시가총액 비율(USD/현지)로 달러 환산. 일본: 엔화 백만"),
        ("수익률", "한국: 1M/3M/6M 원본 컬럼, 1Y는 241거래일 전 종가 대비(가격수익률). 그 외: Datastream PCH#(X, 기간) 가격수익률. 분배금 미포함"),
    ]
    if key == "us":
        rows.append(("미국 액티브 주의", "미국 액티브 ETF는 종목명에 'Active'를 쓰지 않는 경우가 많아(예: Dimensional·Avantis·JPMorgan 계열) 이 규칙으로는 액티브가 과소 집계됩니다."))
    if key == "eu":
        rows.append(("유럽 범위", "아일랜드·룩셈부르크 설정 UCITS ETF. 원본에 대표 상장 RIC 하나씩만 있어 거래소는 그 RIC 기준"))
    if dropped:
        rows.append(("제외", " · ".join(f"{k} {v:,}건" for k, v in dropped.items())))
    return rows


def round_rec(rec):
    out = {}
    for k, v in rec.items():
        if k == "rf":
            continue
        if isinstance(v, float):
            v = round(v, 4 if k == "ter" else 2)
        out[k] = v
    return out


JSON_KEYS = ["t", "n", "iss", "ex", "dom", "scope", "reg", "ac", "sec", "st", "lev", "act", "hd", "bm", "ccy", "ld", "ter",
             "aum", "adv", "r1w", "r1m", "r3m", "r6m", "r1y", "vol", "shp", "gap", "te", "mdd", "dv"]


def write_json(key, recs, stamp, dropped):
    m = MARKETS[key]
    recs = sorted(recs, key=lambda x: -(x["aum"] or 0))
    cols = [k for k in JSON_KEYS if any(r.get(k) is not None for r in recs)]
    rows = [[round_rec(r).get(k) for k in cols] for r in recs]
    payload = {
        "ok": True, "market": key, "label": m["label"], "as_of": stamp, "unit": m["unit"], "adv_unit": m["adv_unit"],
        "count": len(recs), "total_aum": round(sum(r["aum"] or 0 for r in recs), 1),
        "source": f"SavvyDB/ETF_DB/{m['file']}", "dropped": dropped,
        "notes": [f"{k}: {v}" for k, v in notes_for(key, stamp, dropped)],
        "cols": cols, "rows": rows,
    }
    JSON_DIR.mkdir(parents=True, exist_ok=True)
    (JSON_DIR / f"{key}.json").write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    passive = [r for r in recs if not r["act"]]
    return {
        "market": key, "label": m["label"], "as_of": stamp, "unit": m["unit"], "count": len(recs),
        "total_aum": round(sum(r["aum"] or 0 for r in recs), 1),
        "passive_count": len(passive), "passive_aum": round(sum(r["aum"] or 0 for r in passive), 1),
        "top": [r["n"] for r in recs[:3]],
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--markets", nargs="*", default=list(MARKETS))
    ap.add_argument("--no-json", action="store_true")
    args = ap.parse_args()
    today = datetime.now().strftime("%Y%m%d")
    index = []
    idx_path = JSON_DIR / "index.json"
    if idx_path.exists() and set(args.markets) != set(MARKETS):
        index = [x for x in json.loads(idx_path.read_text(encoding="utf-8"))["markets"] if x["market"] not in args.markets]
    for key in args.markets:
        recs, stamp, dropped = load_market(key)
        out = XLSX_DIR / f"ETF_Universe_{MARKETS[key]['label']}_{today}.xlsx"
        build_workbook(key, recs, stamp, dropped, out)
        if not args.no_json:
            index.append(write_json(key, recs, stamp, dropped))
        print(f"{key}: {len(recs):,} ETFs → {out.name} · dropped {dropped}")
    if not args.no_json:
        order = list(MARKETS)
        index.sort(key=lambda x: order.index(x["market"]))
        idx_path.write_text(json.dumps({"generated_at": datetime.now().isoformat(timespec="seconds"), "markets": index},
                                       ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
