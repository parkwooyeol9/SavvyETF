"""ACWI 원자료의 시계열 시트 로더.

지원 시트 (이름은 대소문자·괄호·공백 무시하고 키워드로 찾는다, `_classify`):
    Price(Daily_3Y) / Price      → ri    X(RI) 일간 (현지통화 총수익지수)
    Volume(Daily_3Y) / Volume    → vol   VO 일간 (거래량, 천 주)
    Price_(Monthly_10Y)          → ri_m  X(RI) 월간 10년 (백테스트용)
    EPS · BPS · DPS · SALE       → eps · bps · dps · sal  12M 선행 월간 (보고통화, 10년)


시트 형식 (Datastream DSGRID, Sym=RIC):
    1행  : A1 비어 있음, B1.. 종목 코드(RIC)
    2행  : A2 DSGRID 수식, B2.. 시리즈 설명 ("NVIDIA - TOT RETURN IND", 실패 시 "#ERROR")
    3행~ : A열 날짜, 값
- Price : X(RI) 일간 (현지통화 총수익지수, 배당 재투자)
- EPS/BPS/DPS : EPS1FD12 / BPS1FD12 / DPS1FD12 월간 (12개월 선행, 보고통화)

주의(2026-10-02 파일에서 발견): 월간 시트 아래쪽에 Price 시트를 복사할 때 남은 '일간 Price 값'이
748행 그대로 붙어 있었다. 날짜가 처음으로 거꾸로 가는 행에서 잘라낸다 (`_cut_stale_tail`).

    from Claude_DB.factor.timeseries import load_timeseries
    ts = load_timeseries("Claude_DB/data/ACWI_new.xlsx")   # out/ts_cache_*.pkl 캐시
    ts.ri, ts.eps, ts.bps, ts.dps, ts.quality
"""
from __future__ import annotations

import hashlib
import pickle
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
CACHE_DIR = ROOT / "out"
SHEETS = {"ri": "Price", "eps": "EPS", "bps": "BPS", "dps": "DPS"}     # 구버전 이름 (호환)
KEYS = ("ri", "vol", "ri_m", "eps", "bps", "dps", "sal")


def _classify(name: str) -> str | None:
    """시트 이름 → 키. 'Price(Daily_3Y)'→ri, 'Price_(Monthly_10Y)'→ri_m, 'Volume(...)'→vol, 'SALE'→sal."""
    import re
    n = re.sub(r"[^a-z0-9]", "", name.lower())
    if n.startswith("price") or n.startswith("ri"):
        return "ri_m" if "month" in n else "ri"
    if n.startswith("volume") or n == "vo":
        return "vol"
    for k, keys in (("eps", ("eps",)), ("bps", ("bps",)), ("dps", ("dps",)), ("sal", ("sale", "sal", "sales"))):
        if any(n == x or n.startswith(x) for x in keys):
            return k
    return None
JUMP = 0.40          # 하루 ±40% 이상 = 이상치 후보 (분할·스핀오프 미조정, 데이터 오류)


@dataclass
class TimeSeries:
    ri: pd.DataFrame                 # 일간 총수익지수 (날짜 × 코드)
    eps: pd.DataFrame                # 월간 12M 선행 EPS
    bps: pd.DataFrame
    dps: pd.DataFrame
    vol: pd.DataFrame = field(default_factory=pd.DataFrame)    # 일간 거래량 (천 주)
    ri_m: pd.DataFrame = field(default_factory=pd.DataFrame)   # 월간 총수익지수 (10년)
    sal: pd.DataFrame = field(default_factory=pd.DataFrame)    # 월간 12M 선행 매출
    desc: dict = field(default_factory=dict)         # 코드 → 시리즈 설명
    quality: dict = field(default_factory=dict)      # 점검 결과 요약
    anomalies: pd.DataFrame | None = None            # 일간 급등락 목록


def _cut_stale_tail(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    """날짜가 단조 증가하다 처음 감소하는 지점 이후를 버린다 (이전 내용이 남은 행)."""
    idx = pd.to_datetime(pd.Series(df.index), errors="coerce")
    back = np.where(idx.diff().dt.days.fillna(1).to_numpy() <= 0)[0]
    if len(back):
        cut = int(back[0])
        return df.iloc[:cut], len(df) - cut
    return df, 0


def _read_sheet(ws) -> tuple[pd.DataFrame, list]:
    rows = list(ws.iter_rows(values_only=True))
    codes = [c for c in rows[0][1:]]
    keep = [i for i, c in enumerate(codes) if c not in (None, "")]
    codes = [str(codes[i]) for i in keep]
    desc = [rows[1][1:][i] if len(rows) > 1 else None for i in keep]
    body = [r for r in rows[2:] if r and r[0] is not None]
    idx = [r[0] for r in body]
    vals = [[(r[1:][i] if i < len(r) - 1 else None) for i in keep] for r in body]
    df = pd.DataFrame(vals, index=idx, columns=codes)
    return df, desc


def _to_num(df: pd.DataFrame) -> pd.DataFrame:
    df = df.apply(pd.to_numeric, errors="coerce")
    df.index = pd.to_datetime(df.index)
    df = df[~df.index.duplicated(keep="last")].sort_index()
    return df


def _file_key(xlsx: Path) -> str:
    st = xlsx.stat()
    return hashlib.md5(f"{xlsx.resolve()}|{st.st_size}|{st.st_mtime_ns}".encode()).hexdigest()[:12]


def load_timeseries(xlsx: str | Path, use_cache: bool = True) -> TimeSeries | None:
    """시계열 시트가 없으면 None."""
    xlsx = Path(xlsx).expanduser()
    cache = CACHE_DIR / f"ts_cache_{_file_key(xlsx)}.pkl"
    if use_cache and cache.exists():
        return pickle.loads(cache.read_bytes())
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        k = _classify(name)
        if k and k not in sheets:
            sheets[k] = name
    if "ri" not in sheets:
        return None
    out, q, desc = {k: pd.DataFrame() for k in KEYS}, {"sheets": sheets}, {}
    for key, sheet in sheets.items():
        df, d = _read_sheet(wb[sheet])
        df, dropped = _cut_stale_tail(df)
        df = _to_num(df)
        q[f"{key}_rows"], q[f"{key}_stale_rows_dropped"] = len(df), dropped
        q[f"{key}_first"], q[f"{key}_last"] = str(df.index.min().date()), str(df.index.max().date())
        q[f"{key}_error_series"] = int(sum(str(x).upper().startswith("#ERROR") for x in d))
        q[f"{key}_all_nan"] = int(df.isna().all().sum())
        if key in ("ri", "ri_m"):
            if key == "ri":
                desc = dict(zip(df.columns, d))
            zeros = int((df == 0).sum().sum())
            df = df.mask(df <= 0)                         # 0·음수 지수 = 결측 (BMPS 2023년 등)
            q[f"{key}_zero_values"] = zeros
        if key in ("eps", "bps", "dps", "sal", "ri_m") and len(df) > 1:
            gap = pd.Series(df.index).diff().dt.days.median()
            q[f"{key}_freq_days"] = float(gap)
            if gap < 20:                                    # 월간이어야 할 시트가 일간 → 잘못 붙은 시트
                q[f"{key}_warning"] = "월간 시트인데 날짜 간격이 일간"
        out[key] = df
    ri = out["ri"]
    r = ri.pct_change(fill_method=None)
    big = r.abs().stack()
    big = big[big >= JUMP]
    an = big.rename("ret").reset_index()
    an.columns = ["date", "code", "ret"]
    an["ret"] = r.stack().reindex(pd.MultiIndex.from_frame(an[["date", "code"]])).to_numpy()
    an["last_day"] = an["date"] == ri.index.max()
    q["ri_jumps"] = int(len(an))
    q["ri_jump_stocks"] = int(an["code"].nunique())
    q["ri_late_start"] = int((ri.apply(lambda s: s.first_valid_index()) > ri.index[0] + pd.Timedelta(days=31)).sum())
    if not out["vol"].empty:
        v = out["vol"].reindex(columns=ri.columns)
        q["vol_all_nan"] = int(v.isna().all().sum())
        q["vol_zero_last20"] = int((v.iloc[-20:].fillna(0) == 0).all().sum())
    ts = TimeSeries(ri=ri, eps=out["eps"], bps=out["bps"], dps=out["dps"], vol=out["vol"], ri_m=out["ri_m"],
                    sal=out["sal"], desc=desc, quality=q, anomalies=an)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    for old in CACHE_DIR.glob("ts_cache_*.pkl"):
        old.unlink()
    cache.write_bytes(pickle.dumps(ts))
    return ts
