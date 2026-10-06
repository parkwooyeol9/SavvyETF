"""엑셀 → 팩터 계산 → DB/CSV/JSON/웹 데이터 생성.

사용 (저장소 루트에서):
    python -m Claude_DB.factor.build                                  # data/ACWI_raw.xlsx 사용
    python -m Claude_DB.factor.build --xlsx ~/Downloads/새파일.xlsx --as-of 2026-10-01
    python -m Claude_DB.factor.build --neutral sector                  # 섹터 중립만 (원본 방식)
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd

from . import config as C
from . import engine as E

ROOT = Path(__file__).resolve().parents[1]

# DB/CSV 에 남길 열 (순서 = 화면 기본 순서)
STOCK_COLS = [
    "rank", "code", "name", "sector", "sector_ko", "industry", "country", "region", "currency", "asset_cat",
    "wgt", "ffcap_usd_mn", "mcap_usd_mn", "price", "price_local", "px_usd", "px_unit", "px_unit_fixed",
    "mv_local", "report_ccy", "fx_k", "fx_fixed", "fx_method",
    "eps_ntm_raw", "bps_ntm_raw", "dps_ntm_raw", "ni_usd_mn", "sales_usd_mn", "ret_1m", "ret_3m", "ret_12m", "mom_12_1",
    "eps_ntm", "bps_ntm", "dps_ntm", "ni_ntm", "sales_ntm", "pe", "pb", "ep", "bp", "roe", "dy", "fcfy", "de", "vol",
    "eps_cy25", "eps_cy26", "eps_cy27", "rev_cy25", "rev_cy26", "rev_cy27",
    "eps_g26", "eps_g", "sales_g26", "sales_g", "eps_flag", "dps_g", "ni_ntm_chg", "sales_ntm_chg",
]
Z_COLS = [f"z_{f}" for fk in C.FACTOR_KEYS for f, _, _ in C.FACTORS[fk]["descriptors"]]
SCORE_COLS = ([f"f_{k}" for k in C.FACTOR_KEYS] + [f"p_{k}" for k in C.FACTOR_KEYS]
              + ["composite", "pctile", "sector_rank", "coverage"]
              + [f"o_{k}" for k in C.FACTOR_KEYS] + ["o_composite", "o_rank", "xl_composite", "xl_rank", "cmp_rank",
                                                     "excel_row"])

# 웹 data.js 에 넣을 짧은 키
WEB_MAP = {
    "code": "c", "name": "n", "sector_ko": "s", "industry": "ind", "country": "ct", "region": "rg",
    "currency": "cur", "asset_cat": "ac", "wgt": "w", "mcap_usd_mn": "cap", "price_local": "px",
    "px_usd": "pxu", "px_unit_fixed": "pxf",
    "ret_1m": "r1", "ret_3m": "r3", "ret_12m": "r12", "mom_12_1": "mom",
    "pe": "pe", "pb": "pb", "roe": "roe", "dy": "dy", "fcfy": "fcf", "de": "de", "vol": "vol",
    "eps_g26": "eg26", "eps_g": "eg27", "sales_g26": "sg26", "sales_g": "sg27", "eps_flag": "ef",
    "dps_g": "dg", "ni_usd_mn": "ni", "sales_usd_mn": "sal", "ni_ntm_chg": "nic",
    "ffcap_usd_mn": "ffc", "report_ccy": "rc", "fx_method": "fxm", "eps_ntm": "eps",
    "f_value": "fv", "f_size": "fs", "f_dividend": "fd", "f_growth": "fg", "f_momentum": "fm", "f_quality": "fq",
    "composite": "z", "rank": "rk", "pctile": "pc", "sector_rank": "srk", "coverage": "cov",
    "o_composite": "oz", "cmp_rank": "ork",
}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, (float, np.floating)):
        if math.isnan(v) or math.isinf(v):
            return None
        return float(f"{float(v):.6g}")
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    return v


def excel_fill(xlsx: Path, sheet: str) -> dict:
    """사용자 엑셀에서 계산 열(W·X·BE·BL)의 수식이 몇 행까지 채워져 있는지."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True)
    ws = wb[sheet]
    cols = {"W": 23, "X": 24, "BE": 57, "BL": 64}
    cnt = dict.fromkeys(cols, 0)
    for row in ws.iter_rows(min_row=C.FIRST_ROW, max_row=C.LAST_ROW, max_col=64, values_only=True):
        for k, i in cols.items():
            if i - 1 < len(row) and row[i - 1] not in (None, ""):
                cnt[k] += 1
    wb.close()
    return cnt


def build(xlsx: Path, as_of: str, neutral: str | None = None) -> dict:
    if neutral:
        C.NEUTRALIZE = neutral
    d, stats, excluded, raw = E.run(str(xlsx))
    usd, sheet = d.attrs.get("usd", {}), d.attrs.get("sheet")
    # 비교 기준: 엑셀에 순위가 (거의) 다 계산돼 있으면 그 캐시값, 아니면 원본 로직 전 행 재현값
    xl_ok = d["xl_rank"].notna().mean() > 0.9 and d["xl_rank"].nunique() > 0.9 * len(d)
    d["cmp_rank"] = (d["xl_composite"] if xl_ok else d["o_composite"]).rank(ascending=False, method="min")
    cmp_src = "엑셀 캐시 순위(AJ열)" if xl_ok else "원본 로직 전 행 재현"
    fill = excel_fill(xlsx, sheet)
    d = d.sort_values("rank").reset_index(drop=True)
    out = ROOT / "out"
    out.mkdir(exist_ok=True)

    extra = [c for c in ("eps_rev_3m", "mom_12_1_abs", "mom_12_1_px", "vol_proxy", "data_flag") if c in d]
    stocks = d[[c for c in STOCK_COLS + extra + Z_COLS + SCORE_COLS if c in d]].copy()   # 시계열 없는 파일은 일부 열 없음
    stocks.to_csv(out / "stocks.csv", index=False, encoding="utf-8-sig")
    stats.to_csv(out / "sector_stats.csv", index=False, encoding="utf-8-sig")
    excluded[["code", "name", "wgt", "reason"]].to_csv(out / "excluded.csv", index=False, encoding="utf-8-sig")

    # 집계
    F = [f"f_{k}" for k in C.FACTOR_KEYS]
    by_sector = (d.groupby("sector_ko").agg(n=("code", "size"), wgt=("wgt", "sum"),
                                           pe=("pe", "median"), pb=("pb", "median"), roe=("roe", "median"),
                                           dy=("dy", "median"), eps_g=("eps_g", "median"))
                 .reset_index())
    by_country = (d.groupby(["country", "region"]).agg(n=("code", "size"), wgt=("wgt", "sum"),
                                                     comp=("composite", "mean"), pe=("pe", "median"))
                  .reset_index().sort_values("wgt", ascending=False))
    corr = d[F].rank().corr().round(3)                      # Spearman (scipy 없이)

    meta = {
        "as_of": as_of,
        "built_at": dt.datetime.now().isoformat(timespec="seconds"),
        "source": xlsx.name,
        "n_universe": int(len(d)),
        "n_excluded": int(len(excluded)),
        "usd_scale_mn": round(d.attrs.get("usd_scale", float("nan")), 1) if d.attrs.get("usd_scale") else None,
        "neutralize": C.NEUTRALIZE,
        "winsor": list(C.WINSOR),
        "z_clip": C.Z_CLIP,
        "weights": C.DEFAULT_WEIGHTS,
        "spearman_vs_original": round(float(d["composite"].rank().corr(d["cmp_rank"].rank(ascending=False))), 3),
        "compare_source": cmp_src,
        "excel_fill": fill,
        "original_scored_in_excel": min(fill.values()) if fill else None,
        "sheet": sheet,
        "usd": usd,
        "n_px_unit_fixed": int(d["px_unit_fixed"].sum()),
        "n_fx_fixed": int(d["fx_fixed"].sum()),
        "fx_asof": C.FX_ASOF,
        "timeseries": bool(d.attrs.get("has_timeseries")),
        "ts_quality": E.TS_EXTRAS["ts"].quality if E.TS_EXTRAS.get("ts") is not None else None,
    }
    factors = {k: {"label": v["label"],
                   "descriptors": [{"field": f, "dir": s, "label": l} for f, s, l in v["descriptors"]]}
               for k, v in C.FACTORS.items()}

    # SQLite
    db = out / "acwi_factor.db"
    if db.exists():
        db.unlink()
    with sqlite3.connect(db) as con:
        stocks.to_sql("stocks", con, index=False)
        stats.to_sql("sector_stats", con, index=False)
        excluded[["code", "name", "wgt", "reason"]].to_sql("excluded", con, index=False)
        by_sector.to_sql("agg_sector", con, index=False)
        by_country.to_sql("agg_country", con, index=False)
        pd.DataFrame([{"key": k, "value": json.dumps(v, ensure_ascii=False)} for k, v in meta.items()]).to_sql("meta", con, index=False)
        pd.DataFrame([{"factor": k, "label": v["label"], "descriptor": x["field"], "direction": x["dir"], "desc": x["label"]}
                      for k, v in factors.items() for x in v["descriptors"]]).to_sql("factor_defs", con, index=False)
        con.execute("CREATE UNIQUE INDEX ix_stocks_code ON stocks(code)")
        con.execute("CREATE INDEX ix_stocks_sector ON stocks(sector)")

    summary = {"meta": meta, "factors": factors, "corr": corr.to_dict(),
               "by_sector": by_sector.round(3).to_dict("records"),
               "by_country": by_country.round(3).to_dict("records")}
    (out / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=_clean), encoding="utf-8")

    # 웹 데이터
    wm = {k: v for k, v in (WEB_MAP | {"eps_rev_3m": "er3", "data_flag": "flag"}).items() if k in d}
    rows = [{wm[k]: _clean(r[k]) for k in wm} for r in d[list(wm)].to_dict("records")]
    from .excel_v2 import NOTES
    web = {"meta": meta, "factors": factors, "corr": corr.values.tolist(),
           "notes": [dict(zip(["sev", "where", "issue", "impact", "fix"], n)) for n in NOTES],
           "stocks": rows, "excluded": excluded[["code", "name", "reason"]].fillna("").to_dict("records")}
    js = "window.ACWI = " + json.dumps(web, ensure_ascii=False, separators=(",", ":"), default=_clean) + ";\n"
    (ROOT / "web" / "data.js").write_text(js, encoding="utf-8")
    # 서버 없이 한 파일로 열 수 있는 번들 (data.js 를 인라인)
    idx = ROOT / "web" / "index.html"
    html = idx.read_text(encoding="utf-8") if idx.exists() else ""
    tag = '<script src="data.js"></script>'
    if tag in html:
        (out / "acwi_dashboard_standalone.html").write_text(html.replace(tag, "<script>" + js + "</script>"), encoding="utf-8")
        # 아티팩트(웹 게시)용 조각: doctype/html/head/body 래퍼 제거
        frag = html.replace(tag, "<script>" + js + "</script>")
        for t in ("<!doctype html>\n", '<html lang="ko">\n', "<head>\n", "</head>\n", "<body>\n", "</body>\n", "</html>\n"):
            frag = frag.replace(t, "", 1)
        (out / "acwi_dashboard_fragment.html").write_text(frag, encoding="utf-8")
    return meta


def main():
    ap = argparse.ArgumentParser(description="ACWI 팩터 DB 빌드")
    ap.add_argument("--xlsx", default=str(ROOT / "data" / "ACWI_raw.xlsx"))
    ap.add_argument("--as-of", default=dt.date.today().isoformat())
    ap.add_argument("--neutral", choices=["sector_region", "sector"], default=None)
    a = ap.parse_args()
    meta = build(Path(a.xlsx).expanduser(), a.as_of, a.neutral)
    print(json.dumps(meta, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
