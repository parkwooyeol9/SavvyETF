"""팩터 계산 엔진.

흐름: load_raw → clean → descriptors → standardize(섹터중립 Z) → factor_scores → composite
     + original_scores: 사용자 엑셀 로직을 '모든 행에 채워 넣었을 때'의 결과 (비교용)
"""
from __future__ import annotations

import math
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from . import config as C


# ───────────────────────── 1. 로드 ─────────────────────────
def find_sheet(names: list[str]) -> str:
    for s in C.SHEETS:
        if s in names:
            return s
    raise SystemExit(f"원자료 시트를 찾지 못했습니다. 시트 이름이 {C.SHEETS} 중 하나여야 합니다 (현재: {names})")


def _hdr_text(rows: list, j: int) -> str:
    return " | ".join(str(r[j]).strip() for r in rows if j < len(r) and r[j] not in (None, ""))


def detect_columns(header_rows: list) -> tuple[dict, dict]:
    """헤더(1~3행)에서 Datastream/Refinitiv 요청식을 읽어 필드 → 열 번호(0-based)를 찾는다.

    열 위치가 바뀌거나 열이 추가돼도 요청식이 같으면 따라간다. 반환: (colmap, info)
    info: years(EPS/매출 추정 연도), found(헤더로 찾은 필드), fallback(config 고정 열로 대체한 필드)
    """
    import re
    width = max(len(r) for r in header_rows)
    texts = {j: _hdr_text(header_rows, j) for j in range(width)}
    r3 = {j: str(header_rows[2][j]).strip() if j < len(header_rows[2]) and header_rows[2][j] else "" for j in range(width)}
    colmap, found = {}, set()

    def first(pred, exclude=()):
        for j in range(width):
            if j in exclude:
                continue
            if pred(texts[j], r3[j]):
                return j
        return None

    U = lambda t: t.upper()
    rules = {
        "wgt": lambda t, h: h.lower() == "wgt",
        "code": lambda t, h: h in ("코드", "Code", "CODE"),
        "price": lambda t, h: re.fullmatch(r"X(\(P\))?(~U\$)?", h.replace(" ", "")) is not None,
        "mv_local": lambda t, h: ("(MV)" in U(h) or h.upper() in ("MV", "MV~U$") or "MARKET VALUE" in U(h)) and "PCH" not in U(h),
        "ret_1m": lambda t, h: U(h).startswith("PCH#(X") and re.search(r"\b1M\)", U(h)) and "FD12" not in U(h),
        "ret_3m": lambda t, h: U(h).startswith("PCH#(X") and re.search(r"\b3M\)", U(h)) and "FD12" not in U(h),
        "ret_12m": lambda t, h: U(h).startswith("PCH#(X") and re.search(r"\b12M\)", U(h)) and "FD12" not in U(h),
        "eps_ntm": lambda t, h: ("EPS1FD12" in U(h) or "FORWARD EPS" in U(h)) and "PCH" not in U(h),
        "bps_ntm": lambda t, h: ("BPS1FD12" in U(h) or "FORWARD BPS" in U(h)) and "PCH" not in U(h),
        "roe_ntm": lambda t, h: "FORWARD ROE" in U(h) or "ROE1FD12" in U(h),
        "ni_ntm": lambda t, h: ("INC1FD12" in U(h) or "FORWARD INC" in U(h)) and "PCH" not in U(h),
        "dps_ntm": lambda t, h: ("DPS1FD12" in U(h) or "FORWARD DPS" in U(h)) and "PCH" not in U(h),
        "sales_ntm": lambda t, h: ("SAL1FD12" in U(h) or "FORWARD SAL" in U(h)) and "PCH" not in U(h),
        "ni_ntm_chg": lambda t, h: "PCH" in U(h) and "INC1FD12" in U(h),
        "dps_ntm_chg": lambda t, h: "PCH" in U(h) and "DPS1FD12" in U(h),
        "sales_ntm_chg": lambda t, h: "PCH" in U(h) and "SAL1FD12" in U(h),
        "ric": lambda t, h: U(h) == "RIC",
        "isin": lambda t, h: "ISIN" in U(h),
        "de": lambda t, h: "DEBT % COMMON" in U(h) or "WC08231" in U(h),
        "px_sd_1y": lambda t, h: "SDN#" in U(h) or "VOLATIL" in U(h) or "STDEV" in U(h),
        "fcf_yield": lambda t, h: "FREE CASH FLOW" in U(h),
        "name": lambda t, h: "TR.COMMONNAME" in U(t),
        "industry": lambda t, h: "TR.GICSINDUSTRY" in U(t) and "GROUP" not in U(t),
        "sector": lambda t, h: "TR.GICSSECTOR" in U(t),
        "asset_cat": lambda t, h: "TR.ASSETCATEGORY" in U(t),
        "hq_country": lambda t, h: "ULTIMATEPARENTCOUNTRYHQ" in U(t).replace(" ", "") or "TR.HQCOUNTRY" in U(t),
        "country_src": lambda t, h: any(k in U(t) for k in ("MSCI COUNTRY", "TR.HQCOUNTRYCODE", "TR.EXCHANGECOUNTRYCODE",
                                                            "COUNTRY CODE", "GEOGN", "GGISN")),
    }
    for f, pred in rules.items():
        j = first(pred, exclude=set(colmap.values()))
        if j is not None:
            colmap[f], _ = j, found.add(f)
    # 추정 연도 열: TR.EPSMeanEstimate(Period=CYyyyy) / TR.RevenueMean(Period=CYyyyy) → 연도순 y0,y1,y2
    years = {}
    for kind, key, slots in (("eps", "EPSMEANESTIMATE", ("eps_cy25", "eps_cy26", "eps_cy27")),
                             ("rev", "REVENUEMEAN", ("rev_cy25", "rev_cy26", "rev_cy27"))):
        cols = []
        for j in range(width):
            m = re.search(key + r"\(PERIOD=(?:CY|FY)(\d{4})", U(texts[j]).replace(" ", ""))
            if m:
                cols.append((int(m.group(1)), j))
        cols = sorted(cols)[-3:]
        if len(cols) == 3:
            for (yr, j), slot in zip(cols, slots):
                colmap[slot] = j
                found.add(slot)
            years[kind] = [y for y, _ in cols]
    # 헤더로 못 찾은 필드는 config 고정 열
    fallback = []
    for col, f in C.RAW_COLUMNS.items():
        if f not in colmap:
            colmap[f] = column_index_from_string(col) - 1
            if f not in ("xl_composite", "xl_rank"):
                fallback.append(f)
    hdr = {f: r3.get(j, "") or texts.get(j, "") for f, j in colmap.items()}
    return colmap, {"years": years, "found": sorted(found), "fallback": fallback, "header": hdr}


def vol_kind(header: str) -> str:
    """U열 정의 판별: 'price_level' = SDN#(X,1Y) (주가 수준의 표준편차, 가격 단위), 'return' = 수익률 변동성."""
    import re
    h = (header or "").upper().replace(" ", "")
    if re.search(r"SDN#\(X(\(P\))?(~U\$)?,", h) or h.startswith("SDN#(X1Y") or re.fullmatch(r"SDN#\(X(~U\$)?1Y\)", h):
        return "price_level"
    return "return"


def load_raw(xlsx_path: str) -> pd.DataFrame:
    """UNIVERSE 시트의 원자료 열을 값(data_only)으로 읽는다.

    - 열 위치는 헤더 요청식으로 찾는다(detect_columns). 못 찾으면 config.RAW_COLUMNS 고정 열.
    - 데이터 행은 4행부터 코드(B)가 연속 50행 비면 끝으로 본다(config.LAST_ROW 제한 없음).
    - 헤더에 '~U$' 가 있으면 그 열은 USD. attrs["usd"] = {"price", "mv", "fund"}
    """
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    sheet = find_sheet(wb.sheetnames)
    ws = wb[sheet]
    header_rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=3, values_only=True)]
    colmap, info = detect_columns(header_rows)
    h = info["header"]
    usd = {"price": "~U$" in h.get("price", ""), "mv": "~U$" in h.get("mv_local", ""),
           "fund": all("~U$" in h.get(f, "") for f in ("eps_ntm", "bps_ntm", "dps_ntm"))}
    rows, blank, ci = [], 0, colmap["code"]
    for r in ws.iter_rows(min_row=C.FIRST_ROW, values_only=True):
        if ci < len(r) and r[ci] not in (None, ""):
            blank = 0
        else:
            blank += 1
            if blank >= 50:
                break
        rows.append({f: (r[i] if i < len(r) else None) for f, i in colmap.items()})
    wb.close()
    df = pd.DataFrame(rows)
    df["excel_row"] = range(C.FIRST_ROW, C.FIRST_ROW + len(df))
    text_cols = {"code", "ric", "isin", "name", "industry", "sector", "asset_cat", "country_src", "hq_country"}
    for c in df.columns:
        if c not in text_cols and c != "excel_row":
            df[c] = pd.to_numeric(df[c], errors="coerce")
    df.attrs.update(usd=usd, sheet=sheet, colmap=colmap, col_info=info,
                    vol_kind=vol_kind(h.get("px_sd_1y", "")), n_rows=int(df["code"].notna().sum()))
    return df


# ───────────────────────── 2. 정제 ─────────────────────────
def _country(code: str, isin: str, name: str) -> str:
    if code in C.CODE_COUNTRY_OVERRIDE:
        return C.CODE_COUNTRY_OVERRIDE[code]
    sfx = code.split(".")[-1].split("^")[0] if "." in code else ""
    country = C.SUFFIX_COUNTRY.get(sfx, ("기타", ""))[0]
    if sfx == "HK":
        pre = (isin or "")[:2]
        if pre == "HK" or any(h.lower() in (name or "").lower() for h in C.HK_DM_NAME_HINTS):
            country = "홍콩"
    return country


def _currency(code: str) -> str:
    sfx = code.split(".")[-1].split("^")[0] if "." in code else ""
    return C.SUFFIX_COUNTRY.get(sfx, ("", ""))[1]


def clean(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """분석 대상만 남긴다. (반환: 대상, 제외목록)"""
    d = df[df["code"].notna()].copy()
    reason = pd.Series("", index=d.index)
    reason[d["wgt"].fillna(0) <= 0] = "ACWI 비중 0 (편출·거래정지·선물)"
    reason[(reason == "") & (d["sector"].isna() | (d["sector"] == "NULL"))] = "섹터 없음 (현금성 펀드 등)"
    reason[(reason == "") & d["price"].isna()] = "가격 없음"
    excluded = d[reason != ""].assign(reason=reason[reason != ""])
    d = d[reason == ""].copy()

    d["country"] = [_country(c, i, n) for c, i, n in zip(d["code"], d["isin"], d["name"])]
    d["currency"] = [_currency(c) for c in d["code"]]
    d["region"] = np.where(d["country"] == "미국", "미국",
                           np.where(d["country"].isin(C.DM_COUNTRIES), "선진(미국 외)", "신흥"))
    d["name"] = d["name"].where(d["name"].notna() & (d["name"] != "NULL"), d["code"])
    d["sector_ko"] = d["sector"].map(C.SECTOR_KO).fillna(d["sector"])
    d["is_adr"] = d["asset_cat"].eq("American Depository Receipt")

    # 유동시총(USD): ACWI 비중 × (미국 종목의 MV/비중 중앙값). MV 가 현지통화라서 비중으로 환산한다.
    us = d[d["code"].str.contains(r"\.(?:N|OQ)$", regex=True) & (d["wgt"] > 0) & d["mv_local"].gt(0)]
    k = float((us["mv_local"] / us["wgt"]).median())
    d["ffcap_usd_mn"] = d["wgt"] * k
    d.attrs["usd_scale"] = k
    return d.reset_index(drop=True), excluded.reset_index(drop=True)


# ───────────────────────── 2b. 통화 보정 ─────────────────────────
def _fx_base(cur: str) -> tuple[str, float]:
    """거래통화 → (기준통화, 보조단위 배수). GBp → (GBP, 100)"""
    return C.SUBUNIT.get(cur, (cur, 1))


def fix_currency(d: pd.DataFrame) -> pd.DataFrame:
    """가격(거래통화)과 I/B/E/S 주당 추정치(보고통화)의 통화 불일치를 찾아 거래통화로 맞춘다.

    k = 보고통화 1단위가 거래통화(호가 단위) 몇 단위인지. 후보 k 중 P/E·P/B 가 섹터×지역 중앙값에
    가장 가까워지는 것을 고르되, 개선폭이 FX_FIX_MIN_GAIN(log) 이상일 때만 적용한다.
    """
    d = d.copy()
    base = d["currency"].map(lambda c: _fx_base(c)[0])
    sub = d["currency"].map(lambda c: _fx_base(c)[1]).astype(float)
    fx_t = base.map(C.FX_PER_USD)
    d["mcap_usd_mn"] = d["mv_local"] / fx_t
    pe0 = (d["price"] / d["eps_ntm"]).where(d["eps_ntm"] > 0)
    pb0 = (d["price"] / d["bps_ntm"]).where(d["bps_ntm"] > 0)
    region = np.where(d["country"] == "미국", "미국", np.where(d["country"].isin(C.DM_COUNTRIES), "DM", "EM"))
    grp = d["sector"] + "|" + region
    ok = pe0.between(3, 80) & pb0.between(0.3, 20)
    med_pe = pe0.where(ok).groupby(grp).median()
    med_pb = pb0.where(ok).groupby(grp).median()
    mpe, mpb = grp.map(med_pe), grp.map(med_pb)

    def dist(pe, pb, i):
        parts = []
        if pd.notna(pe) and pd.notna(mpe[i]):
            parts.append(abs(math.log(pe / mpe[i])))
        if pd.notna(pb) and pd.notna(mpb[i]):
            parts.append(abs(math.log(pb / mpb[i])))
        return sum(parts) if parts else np.nan, len(parts)

    ks, rep, how = [], [], []
    fcf = d["fcf_yield"]
    for i in d.index:
        cur = d.at[i, "currency"]
        anchor_only = False
        if cur == "USD":
            cands = C.USD_FOREIGN_CANDIDATES.get(d.at[i, "country"], [])
        elif cur in C.REPORT_CANDIDATES:
            cands = C.REPORT_CANDIDATES[cur]
        else:
            cands, anchor_only = C.ANCHOR_ONLY_CANDIDATES.get(cur, []), True
        # (보고통화, k) 후보. 보조단위 시장은 '보조단위 없이 호가' 경우도 시험 (예: 남아공 일부 종목 ZAR 호가)
        opts = []
        for r in cands:
            if r not in C.FX_PER_USD or pd.isna(fx_t[i]):
                continue
            k = fx_t[i] / C.FX_PER_USD[r] * sub[i]
            opts.append((r, k))
            if sub[i] != 1:
                opts.append((r, k / sub[i]))
        opts = [(r, k) for r, k in opts if abs(math.log(k)) >= 0.3]
        best_k, best_r, method = 1.0, base[i], ""
        ep0 = d.at[i, "eps_ntm"] / d.at[i, "price"] * 100 if d.at[i, "price"] else np.nan
        # 1) 앵커: Datastream FCF 수익률(통화 일관)과 E/P 의 거리
        roe_i = d.at[i, "roe_ntm"]
        use_anchor = (opts and pd.notna(fcf[i]) and fcf[i] > 0 and pd.notna(ep0) and ep0 > 0
                      and pd.notna(roe_i) and roe_i >= 3)
        anchor_agrees = False
        if use_anchor:
            da0 = abs(math.log(ep0 / fcf[i]))
            anchor_agrees = da0 < 0.5
            tried = [(abs(math.log(ep0 * k / fcf[i])), k, r) for r, k in opts if 2.5 <= pe0[i] / k <= 150]
            if tried:
                dmin = min(t[0] for t in tried)
                dk, k, r = next(t for t in tried if t[0] <= dmin + 0.4)
                if da0 - dk >= C.FX_FIX_ANCHOR_GAIN:
                    best_k, best_r, method = k, r, "FCF앵커"
        # 2) 앵커로 결론이 안 나면 섹터×지역 중앙값 (앵커 전용 시장은 생략)
        if not method and not anchor_agrees and opts and not anchor_only:
            b0, n0 = dist(pe0[i], pb0[i], i)
            if pd.notna(b0) and n0 > 0:
                tried = []
                for r, k in opts:
                    pe_k = pe0[i] / k if pd.notna(pe0[i]) else np.nan
                    pb_k = pb0[i] / k if pd.notna(pb0[i]) else np.nan
                    if pd.notna(pe_k) and not (2.5 <= pe_k <= 150):
                        continue
                    dk, nk = dist(pe_k, pb_k, i)
                    if nk == n0:
                        tried.append((dk, k, r))
                if tried:
                    dmin = min(t[0] for t in tried)
                    dk, k, r = next(t for t in tried if t[0] <= dmin + 0.4)   # 후보 순서 우선
                    if (b0 - dk) >= C.FX_FIX_MIN_GAIN * (n0 / 2):
                        best_k, best_r, method = k, r, "섹터중앙값"
        ks.append(best_k); rep.append(best_r); how.append(method)
    d["fx_method"] = how
    d["fx_k"] = ks
    d["report_ccy"] = rep
    d["fx_fixed"] = d["fx_k"] != 1.0
    for c in ("eps_ntm", "bps_ntm", "dps_ntm"):
        d[c + "_raw"] = d[c]
        d[c] = d[c] * d["fx_k"]
    # 순이익·매출(백만, 보고통화) → USD
    d["ni_usd_mn"] = d["ni_ntm"] / d["report_ccy"].map(C.FX_PER_USD)
    d["sales_usd_mn"] = d["sales_ntm"] / d["report_ccy"].map(C.FX_PER_USD)
    return d


def to_usd(d: pd.DataFrame, usd: dict) -> pd.DataFrame:
    """Datastream 이 EPS·BPS·DPS·MV 를 USD(~U$)로 준 경우. 주가만 USD 로 맞추면 된다.

    주가가 현지통화(X)면 환율(FX_PER_USD)과 호가단위로 환산한다. 호가단위는 시장 기본값(영국 펜스=100)이
    대부분 맞지만 종목마다 다를 수 있어(영국 일부 GBP·USD 호가), 후보 단위 중 P/E(없으면 P/B)가
    상식 범위에 드는 것을 고른다. X~U$ 로 받으면 이 단계는 건너뛴다.
    """
    d = d.copy()
    base = d["currency"].map(lambda c: _fx_base(c)[0])
    fx = base.map(C.FX_PER_USD)
    d["price_local"] = d["price"]
    units, changed = [], []
    for i in d.index:
        if usd.get("price"):
            units.append(1.0); changed.append(False); continue
        ovr = C.PRICE_CCY_OVERRIDE.get(d.at[i, "code"])
        if ovr and ovr in C.FX_PER_USD and pd.notna(fx[i]):
            # 단위계수로 표현: price / (fx_base × unit) = price / fx_ovr
            units.append(C.FX_PER_USD[ovr] / fx[i]); changed.append(True); continue
        cands = C.PRICE_UNIT_CANDIDATES.get(d.at[i, "currency"], [_fx_base(d.at[i, "currency"])[1]])
        chosen = cands[0]
        eps, bps, px = d.at[i, "eps_ntm"], d.at[i, "bps_ntm"], d.at[i, "price"]
        if len(cands) > 1 and pd.notna(px) and pd.notna(fx[i]):
            def ok(u):
                pu = px / fx[i] / u
                if pd.notna(eps) and eps > 0:
                    return bool(2 <= pu / eps <= 150)
                if pd.notna(bps) and bps > 0:
                    return bool(0.2 <= pu / bps <= 40)
                return None
            if ok(cands[0]) is False:
                alt = [u for u in cands[1:] if ok(u)]
                if alt:
                    chosen = alt[0]
        units.append(float(chosen)); changed.append(chosen != cands[0])
    d["px_unit"] = units
    d["px_unit_fixed"] = changed
    d["px_usd"] = d["price"] if usd.get("price") else d["price"] / fx / d["px_unit"]
    d["mcap_usd_mn"] = d["mv_local"] if usd.get("mv") else d["mv_local"] / fx
    for c in ("eps_ntm", "bps_ntm", "dps_ntm"):
        d[c + "_raw"] = d[c]
    d["fx_k"], d["report_ccy"], d["fx_fixed"], d["fx_method"] = 1.0, "USD", False, ""
    # 순이익·매출: INC/SAL 단위가 시장마다 다르다(백만/십억) → 시총×E/P 로 순이익, 매출/순이익 비율로 매출
    d["ni_usd_mn"] = d["mcap_usd_mn"] * d["eps_ntm"] / d["px_usd"]
    ratio = (d["sales_ntm"] / d["ni_ntm"]).where(d["ni_ntm"] != 0)
    d["sales_usd_mn"] = d["ni_usd_mn"] * ratio
    return d


def convert(d: pd.DataFrame, usd: dict) -> pd.DataFrame:
    """데이터 형태에 따라 통화를 맞춘다. 결과: px_usd(가치평가용 주가), eps/bps/dps(같은 통화), mcap_usd_mn."""
    if usd.get("fund"):
        return to_usd(d, usd)
    d = fix_currency(d)                       # 구버전(현지통화) 파일: 추정치를 거래통화로 보정
    d["price_local"] = d["price"]
    d["px_usd"] = d["price"]                  # 이 경우 E/P 등은 거래통화 기준(단위 무관)
    d["px_unit"], d["px_unit_fixed"] = 1.0, False
    return d


# ───────────────────────── 3. 디스크립터 ─────────────────────────
def _growth(new: pd.Series, old: pd.Series) -> pd.Series:
    g = (new / old - 1) * 100
    return g.where((new > 0) & (old > 0))          # 적자·흑전·적전은 성장률 정의 불가 → 결측


def _turn_flag(new: pd.Series, old: pd.Series) -> pd.Series:
    out = pd.Series("", index=new.index)
    out[(old > 0) & (new < 0)] = "적전"
    out[(old < 0) & (new < 0)] = "적지"
    out[(old < 0) & (new > 0)] = "흑전"
    return out


def vol_pct(sd: pd.Series, price: pd.Series, kind: str) -> pd.Series:
    """변동성을 '연율 %' 근사로 통일.
    price_level: SDN#(X,1Y) ÷ 주가 × 100 (가격 수준 SD 프록시, 기존 방식)
    return     : 수익률 변동성. 값 크기로 단위 판별 — 중앙값 <0.1 일간 소수, <1 연율 소수, <5 일간 %, 그 이상 연율 %
    """
    if kind == "price_level":
        return sd / price * 100
    m = sd.median()
    if m < 0.1:
        return sd * 100 * np.sqrt(252)
    if m < 1:
        return sd * 100
    if m < 5:
        return sd * np.sqrt(252)
    return sd


def descriptors(d: pd.DataFrame) -> pd.DataFrame:
    d = d.copy()
    p = d["px_usd"]                                      # 추정치와 같은 통화의 주가
    d["pe"] = (p / d["eps_ntm"]).where(d["eps_ntm"] > 0)
    d["pb"] = (p / d["bps_ntm"]).where(d["bps_ntm"] > 0)
    d["ep"] = d["eps_ntm"] / p * 100                     # 적자 기업도 음(-)의 E/P 로 비교 가능
    d["bp"] = (d["bps_ntm"] / p * 100).where(d["bps_ntm"] > 0)
    d["fcfy"] = d["fcf_yield"]
    d["ln_mcap"] = np.log(d["mcap_usd_mn"].where(d["mcap_usd_mn"] > 0))
    d["dy"] = d["dps_ntm"] / p * 100
    d["dps_g"] = d["dps_ntm_chg"].where(d["dps_ntm"] > 0)
    d["eps_g26"] = _growth(d["eps_cy26"], d["eps_cy25"])
    d["eps_g"] = _growth(d["eps_cy27"], d["eps_cy26"])
    d["sales_g26"] = _growth(d["rev_cy26"], d["rev_cy25"])
    d["sales_g"] = _growth(d["rev_cy27"], d["rev_cy26"])
    d["eps_flag"] = _turn_flag(d["eps_cy27"], d["eps_cy26"])
    d["mom_12_1"] = ((1 + d["ret_12m"] / 100) / (1 + d["ret_1m"] / 100) - 1) * 100
    d["roe"] = d["roe_ntm"]
    d["de"] = d["de"].where((d["de"] >= 0) & (d["sector"] != "Financials"))
    d["vol"] = vol_pct(d["px_sd_1y"], d["price_local"], d.attrs.get("vol_kind", "price_level"))
    # 시계열(Price=RI 일간, EPS 월간) 시트가 있으면 더 정확한 값으로 교체 (attach_timeseries 가 넣어 둔 열)
    if "ts_vol_1y" in d:
        d["vol_proxy"] = d["vol"]
        d["vol"] = d["ts_vol_1y"].where(d["ts_vol_1y"].notna(), np.nan)
    if "ts_mom_12_1" in d:
        d["mom_12_1_px"] = d["mom_12_1"]
        d["mom_12_1"] = d["ts_mom_12_1"].where(d["ts_mom_12_1"].notna(), d["mom_12_1"])
    if "ts_eps_rev_3m" in d:
        d["eps_rev_3m"] = d["ts_eps_rev_3m"]
    if "ts_mom_12_1" in d:
        # 국가 상대 모멘텀: 현지통화 총수익이라 고인플레 통화(튀르키예·이집트 등)가 모멘텀 상위를 차지하는 문제 →
        # 국가 중앙값 대비 초과(복리)로 바꾼다. 국가 표본 1개(룩셈부르크 등)면 지역 중앙값.
        g = d["country"].where(d.groupby("country")["mom_12_1"].transform("count") >= 2, d["region"])
        med = d.groupby(g)["mom_12_1"].transform("median")
        d["mom_12_1_abs"] = d["mom_12_1"]
        d["mom_12_1"] = ((1 + d["mom_12_1"] / 100) / (1 + med / 100) - 1) * 100
    if "ts_jump_10d" in d:
        # 최근 급변(스핀오프·분할 미조정 의심): 주가와 추정치 기준 시점이 어긋나 밸류·배당 지표를 쓰지 않는다
        bad = d["ts_jump_10d"].fillna(False).astype(bool)
        d.loc[bad, ["ep", "bp", "dy", "fcfy", "pe", "pb"]] = np.nan
        d["data_flag"] = np.where(bad, "최근 10일 ±40% 급변: 밸류·배당 제외", "")
    return d


def attach_timeseries(d: pd.DataFrame, tech: pd.DataFrame, fund: pd.DataFrame) -> pd.DataFrame:
    """technicals.compute / fundamentals_ts.compute 결과 중 팩터 입력에 쓰는 열을 ts_* 로 붙인다."""
    d = d.copy()
    src = {"ts_vol_1y": (tech, "vol_1y"), "ts_mom_12_1": (tech, "mom_12_1"), "ts_eps_rev_3m": (fund, "eps_rev_3m"),
           "ts_jump_10d": (tech, "jump_10d")}
    for col, (df, c) in src.items():
        if df is not None and c in df:
            d[col] = d["code"].map(df[c])
    if "ts_vol_1y" in d and "suspended" in tech:
        d.loc[d["code"].map(tech["suspended"]).fillna(False).astype(bool), "ts_vol_1y"] = np.nan
    return d


# ───────────────────────── 4. 표준화 ─────────────────────────
def winsorize(s: pd.Series, lo=C.WINSOR[0], hi=C.WINSOR[1]) -> pd.Series:
    v = s.dropna()
    if v.empty:
        return s
    return s.clip(v.quantile(lo), v.quantile(hi))


def sector_z(s: pd.Series, group: pd.Series, parent: pd.Series | None = None) -> tuple[pd.Series, pd.DataFrame]:
    """그룹 내 (x-평균)/표준편차. 평균·표준편차는 같은 윈저라이즈 표본에서 계산.

    group  : 중립화 단위 (기본 '섹터|지역')
    parent : group 표본이 MIN_GROUP_N 미만일 때 대신 쓸 상위 단위 (기본 '섹터')
    """
    g_mean, g_std = s.mean(), s.std(ddof=1)
    stats = s.groupby(group).agg(["mean", "std", "count"])
    mu = group.map(stats["mean"])
    sd = group.map(stats["std"])
    n = group.map(stats["count"]).fillna(0)
    if parent is not None:
        pstats = s.groupby(parent).agg(["mean", "std", "count"])
        small = n < C.MIN_GROUP_N
        mu = mu.where(~small, parent.map(pstats["mean"]))
        sd = sd.where(~small, parent.map(pstats["std"]))
        n = n.where(~small, parent.map(pstats["count"]).fillna(0))
    tiny = n < C.MIN_SECTOR_N
    mu = mu.where(~tiny, g_mean)
    sd = sd.where(~tiny, g_std).replace(0, np.nan)
    z = ((s - mu) / sd).clip(-C.Z_CLIP, C.Z_CLIP)
    return z, stats


def zscore(s: pd.Series) -> pd.Series:
    return (s - s.mean()) / s.std(ddof=1)


# ───────────────────────── 5. 팩터 · 종합 ─────────────────────────
def factor_scores(d: pd.DataFrame, weights: dict | None = None) -> tuple[pd.DataFrame, pd.DataFrame]:
    weights = weights or C.DEFAULT_WEIGHTS
    d = d.copy()
    d["grp"] = d["sector"] + "|" + (d["region"] if C.NEUTRALIZE == "sector_region" else "ALL")
    stat_rows = []
    for fk, f in C.FACTORS.items():
        zs = []
        for field, sign, label in f["descriptors"]:
            if field not in d or d[field].notna().sum() == 0:      # 이 파일에 없는 디스크립터는 건너뜀
                continue
            w = winsorize(d[field])
            z, st = sector_z(w, d["grp"], d["sector"])
            d[f"z_{field}"] = z * sign
            zs.append(f"z_{field}")
            for grp, r in st.iterrows():
                sec, reg = grp.split("|")
                stat_rows.append({"factor": fk, "descriptor": field, "label": label, "sector": sec, "region": reg,
                                  "mean": r["mean"], "std": r["std"], "n": int(r["count"]), "direction": sign})
        raw = d[zs].mean(axis=1, skipna=True)                  # 가용 디스크립터 평균
        d[f"cov_{fk}"] = d[zs].notna().sum(axis=1)
        # 섹터 안에서 다시 표준화 → 팩터마다 분산 1 (가중치가 실제 기여도가 되도록)
        fz, _ = sector_z(raw, d["grp"], d["sector"])
        d[f"f_{fk}"] = fz.fillna(0.0)                           # 결측 팩터 = 섹터 중립(0)
    total_w = sum(weights.values())
    d["composite_raw"] = sum(d[f"f_{k}"] * weights[k] for k in C.FACTOR_KEYS) / total_w
    d["composite"] = zscore(d["composite_raw"])
    d["rank"] = d["composite"].rank(ascending=False, method="min").astype(int)
    d["pctile"] = (d["composite"].rank(pct=True) * 100).round(1)
    d["sector_rank"] = d.groupby("sector")["composite"].rank(ascending=False, method="min").astype(int)
    d["coverage"] = sum((d[f"cov_{k}"] > 0).astype(int) for k in C.FACTOR_KEYS)
    for k in C.FACTOR_KEYS:
        d[f"p_{k}"] = (d[f"f_{k}"].rank(pct=True) * 100).round(0)
    return d, pd.DataFrame(stat_rows)


# ───────────────────────── 6. 원본 엑셀 로직 재현 (비교용) ─────────────────────────
def _trimmean(v: np.ndarray, pct=0.1) -> float:
    v = np.sort(v[~np.isnan(v)])
    n = len(v)
    if n == 0:
        return np.nan
    k = int(math.floor(n * pct / 2))          # Excel TRIMMEAN: 양쪽에서 floor(n*pct/2)개씩 제거
    return float(v[k:n - k].mean()) if n - 2 * k > 0 else float(v.mean())


def original_scores(raw: pd.DataFrame) -> pd.DataFrame:
    """엑셀 수식을 2314행까지 모두 채웠을 때의 결과. 단위·정의 문제는 그대로 둔다."""
    d = raw[raw["code"].notna() & raw["sector"].notna()].copy()
    p = d["price"]
    pe = p / d["eps_ntm"]; d["AA"] = pe.where((d["eps_ntm"] > 0) & (pe < 300))
    pb = p / d["bps_ntm"]; d["AB"] = pb.where((d["bps_ntm"] > 0) & (pb < 30))
    w = d["dps_ntm"] / p * 100; d["W"] = w.where(w < 50)
    d["X"] = d["ret_12m"] - d["ret_1m"]
    d["AX"] = _growth(d["eps_cy27"], d["eps_cy26"])
    d["AZ"] = _growth(d["rev_cy27"], d["rev_cy26"])
    d["T_"] = d["de"]
    spec = [  # (z열, 원자료, 평균방식)
        ("BE", "AA", "avg"), ("BF", "AB", "trim"), ("BG", "mv_local", "avg"), ("BH", "W", "trim"),
        ("BI", "dps_ntm_chg", "trim"), ("BJ", "AX", "trim"), ("BK", "AZ", "trim"), ("BL", "X", "trim"),
        ("BM", "T_", "trim"), ("BN", "px_sd_1y", "trim"), ("BO", "fcf_yield", "trim"),
    ]
    # 원자료 열(빈 셀)은 엑셀 산식에서 0 으로 취급된다: (빈칸-평균)/표준편차, TRIMMEAN(IF(..)) 모두 0 포함
    for c in ("mv_local", "dps_ntm_chg", "T_", "px_sd_1y", "fcf_yield"):
        d[c] = d[c].fillna(0)
    for zc, src, how in spec:
        g = d.groupby("sector")[src]
        mu = g.apply(lambda s: s.mean() if how == "avg" else _trimmean(s.to_numpy(dtype=float)))
        sd = g.std(ddof=1)
        d[zc] = (d[src] - d["sector"].map(mu)) / d["sector"].map(sd)
    d.loc[d["T_"] < 0, "BM"] = np.nan
    d["o_value"] = -d[["BE", "BF"]].mean(axis=1)
    d["o_size"] = -d["BG"]
    d["o_dividend"] = d[["BH", "BI"]].mean(axis=1).fillna(0).clip(upper=3)
    d["o_growth"] = d[["BJ", "BK"]].mean(axis=1).fillna(0).clip(upper=3.5)
    d["o_momentum"] = d["BL"].fillna(0).clip(upper=2.5)
    d["o_quality"] = pd.concat([-d["BM"], -d["BN"], d["BO"]], axis=1).mean(axis=1)
    d["o_composite"] = d[[f"o_{k}" for k in C.FACTOR_KEYS]].fillna(0).sum(axis=1) / 6
    d["o_rank"] = d["o_composite"].rank(ascending=False, method="min").astype(int)
    return d[["code"] + [f"o_{k}" for k in C.FACTOR_KEYS] + ["o_composite", "o_rank"]]


TS_EXTRAS: dict = {}      # 마지막 run() 의 시계열 산출물: ts, tech, fund, bench, charts


def run(xlsx_path: str, weights: dict | None = None, ts="auto"):
    """ts: "auto" = 시계열 시트가 있으면 사용, None = 사용 안 함(구버전 결과 재현), 또는 TimeSeries 객체."""
    raw = load_raw(xlsx_path)
    d, excluded = clean(raw)
    scale = d.attrs.get("usd_scale")
    usd = raw.attrs.get("usd", {})
    d = convert(d, usd)
    d.attrs["vol_kind"] = raw.attrs.get("vol_kind", "price_level")
    tech = fund = None
    if isinstance(ts, str) and ts == "auto":
        from .timeseries import load_timeseries
        ts = load_timeseries(xlsx_path)
    if ts is not None:
        from . import fundamentals_ts as FT
        from . import technicals as TE
        w = raw[raw["code"].notna()].drop_duplicates("code").set_index("code")["wgt"]
        tech, bench = TE.compute(ts.ri, w)
        if getattr(ts, "vol", None) is not None and not ts.vol.empty:
            px = d.set_index("code")["px_usd"] if "px_usd" in d else None
            tech = tech.join(TE.liquidity(ts.vol, ts.ri, px))
        d0 = descriptors(d)                                   # 현재 P/E·P/B·배당 (역사적 위치 계산용)
        fund, charts = FT.compute(ts, d0.set_index("code")[["pe", "pb", "dy"]])
        d = attach_timeseries(d, tech, fund)
    d = descriptors(d)
    d, stats = factor_scores(d, weights)
    orig = original_scores(raw)
    d = d.merge(orig, on="code", how="left")
    d.attrs["usd_scale"] = scale
    d.attrs["usd"] = usd
    d.attrs["sheet"] = raw.attrs.get("sheet")
    d.attrs["vol_kind"] = raw.attrs.get("vol_kind")
    d.attrs["col_info"] = raw.attrs.get("col_info")
    TS_EXTRAS.clear()
    if tech is not None:      # DataFrame 을 attrs 에 넣으면 pandas 연산마다 복사·비교돼 느려서 모듈 변수로 보관
        TS_EXTRAS.update(ts=ts, tech=tech, fund=fund, bench=bench, charts=charts)
        d.attrs["has_timeseries"] = True
    return d, stats, excluded, raw
