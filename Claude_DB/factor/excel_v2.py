"""원본 엑셀에 수정된 팩터 산식 시트(FACTOR_v2 · STATS_v2 · PARAM_v2 · MAP_v2 · 검증노트)를 추가한다.

원본 UNIVERSE 시트는 건드리지 않는다. 3행 헤더에 ~U$ 가 있으면(USD 추정치 파일) 그에 맞는 산식을 쓴다. 새 시트는 모두 엑셀 수식이므로
Datastream/Refinitiv 로 원자료를 새로고침하면 수정 점수도 같이 다시 계산된다.

    python -m Claude_DB.factor.excel_v2 --xlsx Claude_DB/data/ACWI_raw.xlsx --out Claude_DB/excel/ACWI_factor_v2.xlsx
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.formula import ArrayFormula

from . import config as C
from . import engine as E
from .xlsx_inject import inject

U = "'UNIVERSE (2)'"
R0, R1 = C.FIRST_ROW, C.LAST_ROW
DESC = [(f, s, fk, lab) for fk in C.FACTOR_KEYS for f, s, lab in C.FACTORS[fk]["descriptors"]]
ND = len(DESC)                     # 12
NF = len(C.FACTOR_KEYS)            # 6

# FACTOR_v2 열 배치
COL_CODE, COL_NAME, COL_SEC, COL_SFX, COL_CTRY, COL_REG, COL_VALID, COL_GRP, COL_CAP, COL_K = range(1, 11)
C_RAW = 11                         # K.. 원값 12개
C_WIN = C_RAW + ND                 # 윈저라이즈 12개
C_Z = C_WIN + ND                   # 방향 반영 Z 12개
C_FRAW = C_Z + ND                  # 팩터 원점수 6개
C_F = C_FRAW + NF                  # 팩터 Z 6개
C_CRAW = C_F + NF
C_COMP, C_RANK, C_PCT, C_SRANK = C_CRAW + 1, C_CRAW + 2, C_CRAW + 3, C_CRAW + 4

HDR = PatternFill("solid", fgColor="1F2937")
SUB = PatternFill("solid", fgColor="E5E7EB")
INP = PatternFill("solid", fgColor="FEF3C7")
WHITE = Font(color="FFFFFF", bold=True)


def u(col: str, r: int) -> str:
    return f"{U}!${col}{r}"


MODE_USD = False                   # build() 에서 원자료 형태에 맞춰 설정


def raw_formula(field: str, r: int) -> str:
    v = f"$G{r}"                    # valid flag
    D, I, J, M, P, F, H, K, T, Uc, V = (u(c, r) for c in "D I J M P F H K T U V".split())
    AO, AP, AR, AS = u("AO", r), u("AP", r), u("AR", r), u("AS", r)
    sec = f"$C{r}"
    k = f"$J{r}"
    if MODE_USD:                    # 추정치=USD, J = 주가 USD 환산계수 → 분모를 USD 주가로
        num = lambda x: x
        den = f"({D}*{k})"
    else:                           # 추정치=보고통화, J = 통화보정계수 k → 분자를 거래통화로
        num = lambda x: f"{x}*{k}"
        den = D
    body = {
        "ep": f'IF(AND(ISNUMBER({I}),ISNUMBER({D}),ISNUMBER({k})),{num(I)}/{den}*100,"")',
        "bp": f'IF(AND(ISNUMBER({J}),ISNUMBER({D}),ISNUMBER({k})),IF({J}>0,{num(J)}/{den}*100,""),"")',
        "fcfy": f'IF(ISNUMBER({V}),{V},"")',
        "ln_mcap": f'IF(ISNUMBER($I{r}),IF($I{r}>0,LN($I{r}),""),"")',
        "dy": f'IF(AND(ISNUMBER({M}),ISNUMBER({D}),ISNUMBER({k})),{num(M)}/{den}*100,"")',
        "dps_g": f'IF(AND(ISNUMBER({P}),ISNUMBER({M})),IF({M}>0,{P},""),"")',
        "eps_g": f'IF(AND(ISNUMBER({AP}),ISNUMBER({AO})),IF(AND({AP}>0,{AO}>0),({AP}/{AO}-1)*100,""),"")',
        "sales_g": f'IF(AND(ISNUMBER({AS}),ISNUMBER({AR})),IF(AND({AS}>0,{AR}>0),({AS}/{AR}-1)*100,""),"")',
        "mom_12_1": f'IF(AND(ISNUMBER({H}),ISNUMBER({F})),((1+{H}/100)/(1+{F}/100)-1)*100,"")',
        "roe": f'IF(ISNUMBER({K}),{K},"")',
        "de": f'IF(ISNUMBER({T}),IF(AND({T}>=0,{sec}<>"Financials"),{T},""),"")',
        "vol": f'IF(AND(ISNUMBER({Uc}),ISNUMBER({D})),{Uc}/{D}*100,"")',
    }[field]
    return f'=IF({v}=1,{body},"")'


def build(src: Path, dst: Path):
    global U, MODE_USD
    # 파이썬 엔진으로 국가 예외표·단위표를 만든다 → 엑셀과 파이썬 분류 일치
    raw = E.load_raw(str(src))
    usd = raw.attrs.get("usd", {})
    U = f"'{raw.attrs['sheet']}'"
    MODE_USD = bool(usd.get("fund"))
    d, _ = E.clean(raw)
    conv = E.convert(d, usd)
    sfx = d["code"].str.split(".").str[-1].str.split("^").str[0]
    by_sfx = sfx.map(lambda s: C.SUFFIX_COUNTRY.get(s, ("기타", ""))[0])
    ovr = d.loc[d["country"] != by_sfx, ["code", "country"]]

    wb = Workbook()
    wb.remove(wb.active)

    # ── MAP_v2 ──  A:E 접미사표 / G:H 국가 예외 / J 선진국 / L: 모드별 종목 예외
    mp = wb.create_sheet("MAP_v2")
    for c, h in zip("ABCDE", ["접미사", "국가", "통화(호가)", "1USD당 환율", "기본 호가단위"]):
        mp[f"{c}1"] = h
    for i, (k, (ct, cur)) in enumerate(C.SUFFIX_COUNTRY.items(), start=2):
        base, sub = C.SUBUNIT.get(cur, (cur, 1))
        unit = C.PRICE_UNIT_CANDIDATES.get(cur, [sub])[0] if MODE_USD else sub
        mp.cell(i, 1, k); mp.cell(i, 2, ct); mp.cell(i, 3, cur)
        mp.cell(i, 4, C.FX_PER_USD.get(base)).fill = INP
        mp.cell(i, 5, unit).fill = INP
    n_sfx = len(C.SUFFIX_COUNTRY) + 1
    mp["G1"], mp["H1"] = "코드 예외", "국가"
    for i, r in enumerate(ovr.itertuples(), start=2):
        mp.cell(i, 7, r.code); mp.cell(i, 8, r.country)
    n_ovr = max(len(ovr) + 1, 2)
    mp["J1"] = "선진국 목록"
    for i, ct in enumerate(sorted(C.DM_COUNTRIES), start=2):
        mp.cell(i, 10, ct)
    n_dm = len(C.DM_COUNTRIES) + 1
    if MODE_USD:
        ex = conv.loc[conv["px_unit_fixed"], ["code", "px_unit", "currency"]]
        mp["L1"], mp["M1"], mp["N1"] = "호가단위 예외 코드", "호가단위", "통화"
        for i, r in enumerate(ex.itertuples(), start=2):
            mp.cell(i, 12, r.code); mp.cell(i, 13, float(r.px_unit)).fill = INP; mp.cell(i, 14, r.currency)
    else:
        ex = conv.loc[conv["fx_fixed"], ["code", "fx_k", "report_ccy", "currency", "fx_method"]]
        for c, h in zip("LMNOP", ["통화보정 코드", "k", "보고통화(추정)", "거래통화", "판정근거"]):
            mp[f"{c}1"] = h
        for i, r in enumerate(ex.itertuples(), start=2):
            mp.cell(i, 12, r.code); mp.cell(i, 13, round(float(r.fx_k), 6)).fill = INP; mp.cell(i, 14, r.report_ccy)
            mp.cell(i, 15, r.currency); mp.cell(i, 16, r.fx_method)
    n_ex = max(len(ex) + 1, 2)
    for c in "ABCDEGHJLMNOP":
        mp[f"{c}1"].font = Font(bold=True)
    mp.column_dimensions["L"].width = 16
    ex_rng = f"MAP_v2!$L$2:$M${n_ex}"
    sfx_rng, ovr_rng, dm_rng = f"MAP_v2!$A$2:$E${n_sfx}", f"MAP_v2!$G$2:$H${n_ovr}", f"MAP_v2!$J$2:$J${n_dm}"

    # ── PARAM_v2 ──
    pr = wb.create_sheet("PARAM_v2")
    pr["A1"] = "파라미터 (노란 칸은 수정 가능)"; pr["A1"].font = Font(bold=True, size=12)
    pr["A3"], pr["B3"] = "팩터", "가중치"
    for i, k in enumerate(C.FACTOR_KEYS):
        pr.cell(4 + i, 1, C.FACTORS[k]["label"])
        c = pr.cell(4 + i, 2, round(C.DEFAULT_WEIGHTS[k], 6)); c.fill = INP
    pr["A10"], pr["B10"] = "합계", "=SUM(B4:B9)"
    W_RNG = "PARAM_v2!$B$4:$B$9"
    pr["A12"], pr["B12"] = "윈저 하위", C.WINSOR[0]
    pr["A13"], pr["B13"] = "윈저 상위", C.WINSOR[1]
    pr["A14"], pr["B14"] = "Z 절단(±)", C.Z_CLIP
    pr["A15"], pr["B15"] = "그룹 최소 표본", C.MIN_GROUP_N
    pr["A16"], pr["B16"] = "섹터 최소 표본", C.MIN_SECTOR_N
    for r in range(12, 17):
        pr.cell(r, 2).fill = INP
    pr["A18"] = "원자료 형태"
    pr["B18"] = ("추정치·시총 USD(~U$), 주가 " + ("USD" if usd.get("price") else "현지통화→환율 환산 (MAP_v2 D·E열)")) if MODE_USD \
        else "추정치 보고통화 → 통화보정 k (MAP_v2 L:P)"
    pr["A20"], pr["B20"], pr["C20"], pr["D20"], pr["E20"], pr["F20"] = "디스크립터", "팩터", "방향", "P_low", "P_high", "설명"
    for j, (f, s, fk, lab) in enumerate(DESC):
        r = 21 + j
        wcol = L(C_RAW + j)
        pr.cell(r, 1, f); pr.cell(r, 2, C.FACTORS[fk]["label"]); pr.cell(r, 3, s)
        pr.cell(r, 4, f"=_xlfn.PERCENTILE.INC(FACTOR_v2!${wcol}${R0}:${wcol}${R1},$B$12)")
        pr.cell(r, 5, f"=_xlfn.PERCENTILE.INC(FACTOR_v2!${wcol}${R0}:${wcol}${R1},$B$13)")
        pr.cell(r, 6, lab)
    for c in ("A3", "B3", "A20", "B20", "C20", "D20", "E20", "F20"):
        pr[c].font = Font(bold=True)
    pr.column_dimensions["A"].width = 16; pr.column_dimensions["F"].width = 34

    # ── STATS_v2 : 키 = 섹터|지역 (33) + 섹터 (11) + ALL ──
    sectors = list(C.SECTOR_KO)
    regions = ["미국", "선진(미국 외)", "신흥"]
    keys = [f"{s}|{r}" for s in sectors for r in regions] + sectors + ["ALL"]
    st = wb.create_sheet("STATS_v2")
    st["A1"] = "그룹 통계 — 행: 섹터|지역, 섹터, 전체 / 열: 항목별 평균·표준편차·N·적용평균·적용표준편차"
    st["A1"].font = Font(bold=True)
    S0 = 4
    S1 = S0 + len(keys) - 1
    st.cell(3, 1, "그룹키"); st.cell(3, 2, "섹터")
    for i, k in enumerate(keys):
        r = S0 + i
        st.cell(r, 1, k)
        st.cell(r, 2, k.split("|")[0] if k != "ALL" else "ALL")
    blocks = [(f"w:{f}", C_WIN + j) for j, (f, *_ ) in enumerate(DESC)] + \
             [(f"f:{k}", C_FRAW + i) for i, k in enumerate(C.FACTOR_KEYS)]
    stat_col = {}
    grp_rng = f"FACTOR_v2!$H${R0}:$H${R1}"
    sec_rng = f"FACTOR_v2!$C${R0}:$C${R1}"
    for b, (name, src_col) in enumerate(blocks):
        c0 = 3 + b * 5
        stat_col[name] = c0
        st.cell(2, c0, name)
        for k, h in enumerate(["평균", "표준편차", "N", "적용평균", "적용SD"]):
            st.cell(3, c0 + k, h)
        x = f"FACTOR_v2!${L(src_col)}${R0}:${L(src_col)}${R1}"
        for i, key in enumerate(keys):
            r = S0 + i
            if key == "ALL":
                cond = f"ISNUMBER({x})"
                st.cell(r, c0, f"=AVERAGE({x})")
                st.cell(r, c0 + 1, f"=_xlfn.STDEV.S({x})")
                st.cell(r, c0 + 2, f"=COUNT({x})")
            else:
                rng = grp_rng if "|" in key else sec_rng
                st.cell(r, c0, f'=IFERROR(AVERAGEIFS({x},{rng},$A{r}),"")')
                ref = f"{L(c0 + 1)}{r}"
                st.cell(r, c0 + 1).value = ArrayFormula(ref, f'=IFERROR(_xlfn.STDEV.S(IF({rng}=$A{r},IF(ISNUMBER({x}),{x}))),"")')
                st.cell(r, c0 + 2, f"=SUMPRODUCT(({rng}=$A{r})*ISNUMBER({x}))")
        m, sd, n = L(c0), L(c0 + 1), L(c0 + 2)
        all_r = S1
        for i, key in enumerate(keys):
            r = S0 + i
            if key == "ALL":
                st.cell(r, c0 + 3, f"={m}{r}"); st.cell(r, c0 + 4, f"={sd}{r}")
                continue
            # 섹터 행: n<최소섹터 → 전체 / 그룹 행: n<최소그룹 → 섹터 행의 적용값
            if "|" in key:
                sr = S0 + len(sectors) * 3 + sectors.index(key.split("|")[0])
                st.cell(r, c0 + 3, f"=IF({n}{r}>=PARAM_v2!$B$15,{m}{r},{L(c0 + 3)}{sr})")
                st.cell(r, c0 + 4, f"=IF({n}{r}>=PARAM_v2!$B$15,{sd}{r},{L(c0 + 4)}{sr})")
            else:
                st.cell(r, c0 + 3, f"=IF({n}{r}>=PARAM_v2!$B$16,{m}{r},{m}${all_r})")
                st.cell(r, c0 + 4, f"=IF({n}{r}>=PARAM_v2!$B$16,{sd}{r},{sd}${all_r})")
    st.column_dimensions["A"].width = 34
    ST_RNG = f"STATS_v2!$A${S0}:${L(3 + len(blocks) * 5)}${S1}"

    # ── FACTOR_v2 ──
    fs = wb.create_sheet("FACTOR_v2")
    fs["A1"] = "ACWI 팩터 v2 — 윈저라이징(1/99%) → 섹터×지역 중립 Z(±3) → 팩터 재표준화 → 가중합 → 전체 Z. 행 번호는 UNIVERSE (2)와 동일"
    fs["A1"].font = Font(bold=True)
    heads = ["코드", "종목명", "섹터", "접미사", "국가", "지역", "유효", "그룹키", "시가총액$mn",
             "주가USD계수" if MODE_USD else "통화보정k"]
    heads += [f"{f}" for f, *_ in DESC] + [f"w_{f}" for f, *_ in DESC] + [f"z_{f}" for f, *_ in DESC]
    heads += [f"raw_{k}" for k in C.FACTOR_KEYS] + [C.FACTORS[k]["label"] for k in C.FACTOR_KEYS]
    heads += ["종합(가중합)", "종합 Z", "순위", "백분위", "섹터내 순위"]
    groups = [(1, "식별"), (C_RAW, "원값 (디스크립터)"), (C_WIN, "윈저라이즈"), (C_Z, "그룹 Z (방향 반영)"),
              (C_FRAW, "팩터 원점수"), (C_F, "팩터 Z (재표준화)"), (C_CRAW, "종합")]
    for c, t in groups:
        fs.cell(2, c, t).font = Font(bold=True)
    for j, h in enumerate(heads, start=1):
        cell = fs.cell(3, j, h); cell.fill = HDR; cell.font = WHITE
        cell.alignment = Alignment(horizontal="center")

    for r in range(R0, R1 + 1):
        code = u("B", r)
        fs.cell(r, COL_CODE, f'=IF({code}="","",{code})')
        fs.cell(r, COL_NAME, f'=IF({u("AK", r)}="","",{u("AK", r)})')
        fs.cell(r, COL_SEC, f'=IF({u("AM", r)}="","",{u("AM", r)})')
        fs.cell(r, COL_SFX, f'=IFERROR(IF(ISERROR(FIND(".",$A{r})),"",LEFT(TRIM(RIGHT(SUBSTITUTE($A{r},".",REPT(" ",60)),60))&"^",FIND("^",TRIM(RIGHT(SUBSTITUTE($A{r},".",REPT(" ",60)),60))&"^")-1)),"")')
        fs.cell(r, COL_CTRY, f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},{ovr_rng},2,FALSE),IFERROR(VLOOKUP($D{r},{sfx_rng},2,FALSE),"기타")))')
        fs.cell(r, COL_REG, f'=IF($E{r}="","",IF($E{r}="미국","미국",IF(COUNTIF({dm_rng},$E{r})>0,"선진(미국 외)","신흥")))')
        fs.cell(r, COL_VALID, f'=IF(AND($A{r}<>"",ISNUMBER({u("A", r)}),$C{r}<>"",$C{r}<>"NULL",ISNUMBER({u("D", r)})),IF({u("A", r)}>0,1,0),0)')
        fs.cell(r, COL_GRP, f'=IF($G{r}=1,$C{r}&"|"&$F{r},"")')
        fxv = f"VLOOKUP($D{r},{sfx_rng},4,FALSE)"
        if MODE_USD:
            mcap = u("E", r) if usd.get("mv") else f'{u("E", r)}/{fxv}'
            unit = f"IFERROR(VLOOKUP($A{r},{ex_rng},2,FALSE),VLOOKUP($D{r},{sfx_rng},5,FALSE))"
            kf = "1" if usd.get("price") else f"1/({fxv}*{unit})"
        else:
            mcap = f'{u("E", r)}/{fxv}'
            kf = f"IFERROR(VLOOKUP($A{r},{ex_rng},2,FALSE),1)"
        fs.cell(r, COL_CAP, f'=IF($G{r}=1,IFERROR({mcap},""),"")')
        fs.cell(r, COL_K, f'=IF($G{r}=1,IFERROR({kf},""),"")')
        for j, (f, s, fk, _) in enumerate(DESC):
            fs.cell(r, C_RAW + j, raw_formula(f, r))
            rc = f"{L(C_RAW + j)}{r}"
            pr_r = 21 + j
            fs.cell(r, C_WIN + j, f'=IF({rc}="","",MIN(MAX({rc},PARAM_v2!$D${pr_r}),PARAM_v2!$E${pr_r}))')
            wc = f"{L(C_WIN + j)}{r}"
            c0 = stat_col[f"w:{f}"]
            mu = f"VLOOKUP($H{r},{ST_RNG},{c0 + 3},FALSE)"
            sd = f"VLOOKUP($H{r},{ST_RNG},{c0 + 4},FALSE)"
            fs.cell(r, C_Z + j, f'=IF({wc}="","",IFERROR(PARAM_v2!$C${pr_r}*MAX(-PARAM_v2!$B$14,MIN(PARAM_v2!$B$14,({wc}-{mu})/{sd})),""))')
        j0 = 0
        for i, k in enumerate(C.FACTOR_KEYS):
            nd = len(C.FACTORS[k]["descriptors"])
            zr = f"{L(C_Z + j0)}{r}:{L(C_Z + j0 + nd - 1)}{r}"
            j0 += nd
            fs.cell(r, C_FRAW + i, f'=IF($G{r}=1,IFERROR(AVERAGE({zr}),""),"")')
            fr = f"{L(C_FRAW + i)}{r}"
            c0 = stat_col[f"f:{k}"]
            mu = f"VLOOKUP($H{r},{ST_RNG},{c0 + 3},FALSE)"
            sd = f"VLOOKUP($H{r},{ST_RNG},{c0 + 4},FALSE)"
            fs.cell(r, C_F + i, f'=IF($G{r}<>1,"",IF({fr}="",0,IFERROR(MAX(-PARAM_v2!$B$14,MIN(PARAM_v2!$B$14,({fr}-{mu})/{sd})),0)))')
        frng = f"{L(C_F)}{r}:{L(C_F + NF - 1)}{r}"
        wsum = "+".join(f"{L(C_F + i)}{r}*PARAM_v2!$B${4 + i}" for i in range(NF))
        fs.cell(r, C_CRAW, f'=IF($G{r}=1,({wsum})/SUM({W_RNG}),"")')
        crc = f"${L(C_CRAW)}${R0}:${L(C_CRAW)}${R1}"
        fs.cell(r, C_COMP, f'=IF($G{r}=1,({L(C_CRAW)}{r}-AVERAGE({crc}))/_xlfn.STDEV.S({crc}),"")')
        cc = f"${L(C_COMP)}${R0}:${L(C_COMP)}${R1}"
        fs.cell(r, C_RANK, f'=IF($G{r}=1,_xlfn.RANK.EQ({L(C_COMP)}{r},{cc},0),"")')
        fs.cell(r, C_PCT, f'=IF($G{r}=1,ROUND(_xlfn.PERCENTRANK.INC({cc},{L(C_COMP)}{r})*100,1),"")')
        fs.cell(r, C_SRANK, f'=IF($G{r}=1,COUNTIFS($C${R0}:$C${R1},$C{r},{cc},">"&{L(C_COMP)}{r})+1,"")')
    fs.freeze_panes = "C4"
    fs.column_dimensions["B"].width = 30
    for j in range(C_RAW, C_SRANK + 1):
        fs.column_dimensions[L(j)].width = 9
        for r in range(R0, R1 + 1):
            fs.cell(r, j).number_format = "0.00"

    # ── 검증노트 ──
    nt = wb.create_sheet("검증노트")
    notes = NOTES
    nt["A1"] = "팩터점수 검증 결과 및 v2 수정 내역 (원본 UNIVERSE (2) 시트는 그대로, 수정 산식은 FACTOR_v2)"; nt["A1"].font = Font(bold=True, size=14)
    nt.append([])
    nt.append(["#", "심각도", "원본 위치", "문제", "영향", "v2 수정"])
    for c in range(1, 7):
        nt.cell(3, c).fill = HDR; nt.cell(3, c).font = WHITE
    for i, row in enumerate(notes, start=1):
        nt.append([i] + list(row))
    for col, w in zip("ABCDEF", (4, 8, 18, 60, 50, 55)):
        nt.column_dimensions[col].width = w
    for row in nt.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    # 원본 패키지(Datastream 메타·조건부서식 보존)에 새 시트만 추가
    inject(src, wb, dst, ["FACTOR_v2", "검증노트", "PARAM_v2", "STATS_v2", "MAP_v2"])


NOTES = [
    ("치명", "W, X, BE:BO", "배당수익률(W)·가격모멘텀(X) 수식이 4~22행(19종목)에만 있음. (9/30 파일은 Z-Score BE~BO 도 19행만, 10/2 파일은 BE~BO 만 전 행으로 채워짐)",
     "W·X 가 비어 2,247종목의 모멘텀 0, 배당은 DPS변화로만 계산. 배당·모멘텀 섹터 통계도 19개 값으로 계산됨",
     "모든 행에 동일 수식 (FACTOR_v2)"),
    ("치명", "E (MV) → BG 사이즈", "Datastream MV 는 현지통화(백만). TSMC 64,312,260(TWD), 삼성전자 KRW",
     "같은 섹터 안에서 원화·엔화 종목이 수천 배 큰 회사로 계산됨. 삼성전자 사이즈 Z -12.2. 섹터 평균·표준편차도 오염",
     "USD 시가총액(10/2 파일은 MV~U$, 구파일은 환율 환산) → ln 변환 후 Z"),
    ("치명", "U (SDN#(X,1Y)) → BN", "SDN#(X,1Y)는 '주가 수준'의 표준편차(가격 단위). 주가와 상관 0.97",
     "삼성전자 77,639 vs 엔비디아 16.3 — 변동성이 아니라 통화·주가 수준을 측정. SK하이닉스 퀄리티 -2.86",
     "SD÷주가(%)로 정규화 (권장: Datastream 에서 일간수익률 변동성으로 교체)"),
    ("치명", "I·J·M (EPS·BPS·DPS) vs D (가격)", "가격은 거래통화, I/B/E/S 추정치는 보고통화. 영국(펜스 vs 달러·파운드), 칠레·인도네시아·멕시코(현지 vs 달러), 북유럽(현지 vs 유로), 중국 ADR(달러 vs 위안)",
     "셸 P/E 714배, HSBC 812배, 인도네시아 AMMN 24만배 / 반대로 YAZG(싱가포르) P/E 1.7배·배당 28.7%로 가치·배당 1위. 65종목 해당(2026-09-30 데이터)",
     "10/2 파일: EPS·BPS·DPS·MV 를 ~U$ 로 받아 해결. 단 주가 X 는 아직 현지통화 → 환율·호가단위로 USD 환산(MAP_v2 D·E, L:M). X~U$ 추가 시 완전 해결"),
    ("높음", "BI·BM·BN·BO, AO~AX 통계", "원자료 빈 셀이 (빈칸-평균)/SD 계산과 TRIMMEAN(IF())·STDEV(IF())에서 0으로 들어감",
     "FCF 수익률 결측 971종목(43%)이 'FCF 0%'로 퀄리티 감점. D/E·DPS변화도 같은 문제", "ISNUMBER 로 결측 제외, 결측은 팩터 평균에서 빠짐"),
    ("높음", "AN~AX vs AN2338~", "평균은 TRIMMEAN 10%(P/E·EV 는 단순평균)인데 표준편차는 절사 없는 전체 STDEV",
     "중심과 척도의 표본이 달라 Z 가 0 중심이 아니고, 이상치가 큰 항목은 SD 가 부풀어 모든 종목 Z 가 0 근처로 눌림",
     "전체 1%/99% 윈저라이징 후 같은 표본으로 평균·SD"),
    ("높음", "AE·AF·AG", "상단만 3 / 3.5 / 2.5 로 절단, 하단은 무제한. 가치·사이즈·퀄리티는 절단 없음",
     "TSLA 모멘텀 -10.0, 가치 -4.2 같은 극단값이 종합점수를 지배. 팩터마다 기준이 다름", "모든 Z 를 ±3 대칭 절단"),
    ("높음", "AC~AH → AI", "디스크립터 Z 평균을 다시 표준화하지 않고 가중합",
     "평균한 팩터는 분산이 1보다 작아(원본 배당 SD 0.60, 퀄리티 0.53) 1/6 동일가중이 실제로는 사이즈·가치 쏠림",
     "팩터 원점수를 그룹 내 다시 Z → 가중합 → 전체 Z"),
    ("중간", "AA·AB (P/E·P/B)", "가격배수를 그대로 Z, 적자·PER>300·PBR>30 은 제외",
     "배수는 오른쪽 꼬리가 길고 비선형. 적자 기업은 가치 점수가 아예 없음", "E/P·B/P(수익률 형태) 사용, 적자는 음의 E/P 로 포함"),
    ("중간", "AH 퀄리티", "D/E·변동성·FCF 수익률로 구성, ROE(K열)는 수집했지만 미사용",
     "FCF 수익률은 가치 지표, 금융업 D/E 는 의미 없음", "퀄리티 = ROE(+)·D/E(-, 금융 제외)·변동성(-) / FCF 수익률은 가치로 이동"),
    ("중간", "섹터 중립화", "글로벌 유니버스를 섹터로만 중립화",
     "수정 후에도 섹터만 중립하면 상위 200 중 80%가 신흥국(소형·저PER). 국가 리스크가 점수에 섞임", "섹터×지역(미국/선진/신흥) 중립, 표본<10 이면 섹터로 대체"),
    ("낮음", "X (12-1M)", "H-F 단순 차감", "복리 기준과 차이 (12M +100%, 1M +20% 면 80%p vs 66.7%)", "(1+R12)/(1+R1)-1"),
    ("낮음", "AJ RANK", "공란·문자 포함 범위에서 RANK", "동률 다수", "유효 종목만 RANK.EQ, 백분위·섹터내 순위 추가"),
    ("참고", "F·G·H 수익률", "현지통화 수익률", "튀르키예·이집트 등 고인플레 국가는 모멘텀·성장이 명목으로 부풀려짐",
     "그룹(신흥) 중립으로 일부 완화. 권장: Datastream USD 수익률(예: PCH#(X(RI)~U$,12M))"),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default="Claude_DB/data/ACWI_raw.xlsx")
    ap.add_argument("--out", default="Claude_DB/excel/ACWI_factor_v2.xlsx")
    a = ap.parse_args()
    build(Path(a.xlsx), Path(a.out))
    print("saved", a.out)


if __name__ == "__main__":
    main()
